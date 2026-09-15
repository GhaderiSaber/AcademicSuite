#!/usr/bin/env python3
"""
Supervisor Sidebar Presentation Preset (persian-defense-presentation-builder v3.6.0)
===================================================================================
Generates publication-grade, supervisor-compliant academic defense presentations (.pptx)
and vector PDFs (.pdf) featuring a persistent right-hand 5-chapter navigation sidebar menu
and dedicated 3-tier hypothesis explanation cards for Chapter 5 deep discussion.

Directly solves the two most common defense objections in Iranian universities:
1. Supervisor mandates their own faculty/department template geometry (e.g. 5-pill sidebar ribbon).
2. Defense committee / supervisor complains: "مطالب مربوط به فصل پنج و تبیین‌ها کم است".

Usage CLI:
    python preset_supervisor_sidebar.py --input presentation_payload.json --output Defense_Presentation.pptx --export-pdf
"""

import os
import sys
import json
import argparse
from typing import Dict, List, Any, Optional

if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Ensure local script imports work
current_dir = os.path.dirname(os.path.abspath(__file__))
if current_dir not in sys.path:
    sys.path.insert(0, current_dir)

from presentation_schema import PALETTES
from rtl_typography import (
    apply_p_rtl,
    apply_text_frame_rtl,
    set_run_font,
    attach_speaker_notes,
    FONT_TITLE,
    FONT_BODY,
    FONT_ENG
)
from layout_engine import add_sidebar_navigation_menu, build_hypothesis_explanation_slide


CHAPTER_NAMES = [
    "فصل اول: کلیات پژوهش",
    "فصل دوم: مبانی و پیشینه",
    "فصل سوم: روش‌شناسی",
    "فصل چهارم: یافته‌های پژوهش",
    "فصل پنجم: بحث و نتیجه‌گیری"
]


class SupervisorSidebarPresentationBuilder:
    """Builder for presentations with persistent right-hand 5-chapter navigation sidebar."""

    def __init__(self, theme_name: str = "academic_navy"):
        self.theme_name = theme_name
        self.palette = PALETTES.get(theme_name, PALETTES["academic_navy"])
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.500)
        self.blank_layout = self.prs.slide_layouts[6]

    def add_sidebar_ribbon(self, slide, active_chapter_idx: int):
        """Adds persistent right-hand 5-pill sidebar navigation."""
        add_sidebar_navigation_menu(
            slide,
            active_index=active_chapter_idx,
            chapters=CHAPTER_NAMES,
            palette=self.palette
        )

    def add_slide_header(self, slide, title: str, subtitle: str = "", left_in: float = 0.60, width_in: float = 9.60):
        """Standard title area on left canvas."""
        t_box = slide.shapes.add_textbox(Inches(left_in), Inches(0.35), Inches(width_in), Inches(1.05))
        tf = t_box.text_frame
        tf.word_wrap = True
        apply_text_frame_rtl(tf)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        
        p_title = tf.paragraphs[0]
        apply_p_rtl(p_title, PP_ALIGN.RIGHT)
        r_title = p_title.add_run()
        set_run_font(r_title, title, FONT_TITLE, 24.0, bold=True, color_rgb=self.palette["primary"])

        if subtitle:
            p_sub = tf.add_paragraph()
            apply_p_rtl(p_sub, PP_ALIGN.RIGHT)
            r_sub = p_sub.add_run()
            set_run_font(r_sub, subtitle, FONT_BODY, 13.5, bold=False, color_rgb=self.palette["text_muted"])

    def create_cover_slide(self, meta: Dict[str, Any], notes: str = ""):
        """Slide 1: High-contrast cover slide."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        
        # Dark navy background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.500))
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.palette["cover_bg"]
        bg.line.fill.background()

        # Center card container
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))
        card.fill.solid()
        card.fill.fore_color.rgb = self.palette["cover_card"]
        card.line.color.rgb = self.palette["accent"]
        card.line.width = Pt(2.0)

        tf = card.text_frame
        tf.word_wrap = True
        apply_text_frame_rtl(tf)

        # University / Degree header
        p0 = tf.paragraphs[0]
        apply_p_rtl(p0, PP_ALIGN.CENTER)
        r0 = p0.add_run()
        uni_str = f"{meta.get('university', 'دانشگاه آزاد اسلامی')} | {meta.get('faculty', 'دانشکده علوم انسانی')}"
        set_run_font(r0, uni_str, FONT_BODY, 14.0, bold=False, color_rgb=self.palette["accent_light"])

        # Title
        p_title = tf.add_paragraph()
        apply_p_rtl(p_title, PP_ALIGN.CENTER)
        r_title = p_title.add_run()
        set_run_font(r_title, meta.get("title", "عنوان پایان‌نامه / رساله"), FONT_TITLE, 25.0, bold=True, color_rgb=self.palette["text_light"])

        # Subtitle
        p_sub = tf.add_paragraph()
        apply_p_rtl(p_sub, PP_ALIGN.CENTER)
        r_sub = p_sub.add_run()
        set_run_font(r_sub, meta.get("subtitle", "جلسه دفاع از پایان‌نامه کارشناسی ارشد"), FONT_BODY, 15.0, bold=False, color_rgb=self.palette["text_muted"])

        # Committee grid text
        p_comm = tf.add_paragraph()
        apply_p_rtl(p_comm, PP_ALIGN.CENTER)
        comm_lines = [
            f"دانشجو: {meta.get('author', '')}",
            f"استاد راهنما: {meta.get('supervisor', '')}",
            f"استاد مشاور: {meta.get('advisor', '')}" if meta.get('advisor') else "",
            f"تاریخ دفاع: {meta.get('defense_date', 'شهریور ۱۴۰۵')}"
        ]
        comm_str = "   |   ".join([c for c in comm_lines if c])
        r_comm = p_comm.add_run()
        set_run_font(r_comm, comm_str, FONT_BODY, 13.5, bold=True, color_rgb=self.palette["text_light"])

        if notes:
            attach_speaker_notes(slide, notes)
        return slide

    def create_content_slide(self, chapter_idx: int, title: str, subtitle: str = "", notes: str = ""):
        """Generic content slide with right sidebar menu."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        
        # Slide Canvas Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.500))
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.palette["bg_slide"]
        bg.line.fill.background()

        # Persistent Sidebar
        self.add_sidebar_ribbon(slide, chapter_idx)

        # Header Title
        self.add_slide_header(slide, title, subtitle)

        if notes:
            attach_speaker_notes(slide, notes)
        return slide

    def create_hypothesis_explanation_slide(
        self,
        title: str,
        subtitle: str,
        finding_text: str,
        mechanism_text: str,
        literature_text: str,
        notes: str = ""
    ):
        """Creates a dedicated 3-tier hypothesis explanation slide for Chapter 5."""
        slide_data = {
            "title": title,
            "subtitle": subtitle,
            "finding_text": finding_text,
            "mechanism_text": mechanism_text,
            "literature_text": literature_text,
            "use_sidebar": True,
            "sidebar_active_index": 4, # Chapter 5
            "speaker_notes": notes
        }
        build_hypothesis_explanation_slide(
            prs=self.prs,
            meta={},
            slide_data=slide_data,
            palette=self.palette,
            slide_num=len(self.prs.slides) + 1,
            total_slides=30
        )
        return self.prs.slides[-1]

    def save(self, output_pptx_path: str):
        """Saves presentation to PPTX."""
        os.makedirs(os.path.dirname(os.path.abspath(output_pptx_path)), exist_ok=True)
        self.prs.save(output_pptx_path)
        print(f"[+] PPTX successfully saved to: {output_pptx_path} ({os.path.getsize(output_pptx_path):,} bytes)")

    def export_pdf(self, pptx_path: str, output_pdf_path: str) -> bool:
        """Headless export to PDF via PowerPoint COM automation."""
        print(f"[+] Exporting to vector PDF via PowerPoint COM: {output_pdf_path}...")
        try:
            import win32com.client
            ppt = win32com.client.Dispatch("PowerPoint.Application")
            presentation = ppt.Presentations.Open(os.path.abspath(pptx_path))
            presentation.SaveAs(os.path.abspath(output_pdf_path), 32) # 32 = ppSaveAsPDF
            presentation.Close()
            ppt.Quit()
            print(f"[+] PDF successfully exported: {output_pdf_path} ({os.path.getsize(output_pdf_path):,} bytes)")
            return True
        except Exception as e:
            print(f"[-] PowerPoint COM PDF export warning: {e}", file=sys.stderr)
            return False


def build_from_payload(payload: Dict[str, Any], output_pptx: str, export_pdf: bool = True, theme_name: str = "academic_navy"):
    """Compiles presentation from structured dictionary payload."""
    builder = SupervisorSidebarPresentationBuilder(theme_name=theme_name)
    meta = payload.get("meta", {})
    slides = payload.get("slides", [])

    # Slide 1: Cover
    builder.create_cover_slide(meta, notes=slides[0].get("speaker_notes", "") if slides else "")

    for s_idx, s in enumerate(slides[1:], start=2):
        layout = s.get("layout", "content")
        ch_idx = s.get("chapter_index", s.get("sidebar_active_index", 0))
        title = s.get("title", "")
        subtitle = s.get("subtitle", "")
        notes = s.get("speaker_notes", "")

        if layout == "hypothesis_explanation":
            builder.create_hypothesis_explanation_slide(
                title=title,
                subtitle=subtitle,
                finding_text=s.get("finding_text", s.get("statistic", "")),
                mechanism_text=s.get("mechanism_text", s.get("interpretation", "")),
                literature_text=s.get("literature_text", s.get("concordance", "")),
                notes=notes
            )
        else:
            builder.create_content_slide(
                chapter_idx=ch_idx,
                title=title,
                subtitle=subtitle,
                notes=notes
            )

    builder.save(output_pptx)
    if export_pdf:
        pdf_path = os.path.splitext(output_pptx)[0] + ".pdf"
        builder.export_pdf(output_pptx, pdf_path)


def main():
    parser = argparse.ArgumentParser(description="Supervisor Sidebar Presentation Builder Preset v3.6.0")
    parser.add_argument("--input", "-i", type=str, help="Path to input presentation payload JSON")
    parser.add_argument("--output", "-o", type=str, default="Defense_Presentation_Supervisor_Format.pptx", help="Path to output PPTX")
    parser.add_argument("--export-pdf", action="store_true", default=True, help="Automatically export to PDF via PowerPoint COM")
    parser.add_argument("--theme", type=str, default="academic_navy", help="Color palette theme name")
    args = parser.parse_args()

    if args.input and os.path.exists(args.input):
        with open(args.input, "r", encoding="utf-8") as f:
            payload = json.load(f)
        build_from_payload(payload, args.output, export_pdf=args.export_pdf, theme_name=args.theme)
    else:
        print("[*] No input payload provided; generating demonstration sidebar presentation...")
        demo_payload = {
            "meta": {
                "title": "ارزیابی مدل ساختاری تروما و رفتارهای پرخطر با میانجی‌گری آگاهی هیجانی و احساس انسجام",
                "subtitle": "جلسه دفاع از پایان‌نامه کارشناسی ارشد روان‌شناسی بالینی",
                "university": "دانشگاه آزاد اسلامی واحد کاشان",
                "faculty": "دانشکده علوم انسانی",
                "author": "مرضیه سینائی",
                "supervisor": "دکتر حمید امیری",
                "advisor": "",
                "defense_date": "شهریور ۱۴۰۵"
            },
            "slides": [
                {"layout": "cover", "speaker_notes": "سلام و عرض احترام خدمت اعضای محترم هیئت داوران..."},
                {
                    "layout": "hypothesis_explanation",
                    "chapter_index": 4,
                    "title": "تبیین فرضیه اول: ترومای دوران کودکی و رفتارهای پرخطر",
                    "subtitle": "تحلیل سه‌سطحی: یافته آماری، مکانیزم نظری و پیشینه همسو",
                    "finding_text": "مسیر مستقیم تروما به رفتارهای پرخطر با ضریب استاندارد β = 0.24 و مقدار p = .005 در سطح ۰.۰۱ تأیید شد.",
                    "mechanism_text": "استرس سمی دوران رشد باعث بیش‌فعالی محور HPA و نقص در مهار تکانه قشر پیش‌پیشانی می‌شود؛ طبق فرضیه خوددرمانی خانتزیان، نوجوان برای تسکین درد عاطفی به رفتارهای پرخطر روی می‌آورد.",
                    "literature_text": "همسو با یافته‌های تئودور و همکاران (2023)، سینک و فاستر (2022) و رفیعی و همکاران (1401).",
                    "speaker_notes": "فرضیه اول با ضریب ۲۴ صدم تایید شد. استرس مزمن دوران کودکی ساختارهای مهار هیجان را در قشر پیش‌پیشانی تضعیف می‌کند."
                }
            ]
        }
        build_from_payload(demo_payload, args.output, export_pdf=args.export_pdf, theme_name=args.theme)


if __name__ == "__main__":
    main()