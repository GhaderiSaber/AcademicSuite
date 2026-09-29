import sys
import docx
from docx.shared import RGBColor
from docx.enum.text import WD_COLOR_INDEX
import os

def replace_in_runs(paragraph, old_text, new_text, track=False):
    full_text = "".join(r.text for r in paragraph.runs)
    if old_text not in full_text:
        return False
        
    start_idx = full_text.find(old_text)
    end_idx = start_idx + len(old_text)
    
    run_offsets = []
    current_len = 0
    for i, r in enumerate(paragraph.runs):
        run_offsets.append((current_len, current_len + len(r.text), i))
        current_len += len(r.text)
        
    affected_runs = []
    for start, end, i in run_offsets:
        if start < end_idx and end > start_idx:
            affected_runs.append(i)
            
    if not affected_runs:
        return False
        
    first_run_idx = affected_runs[0]
    last_run_idx = affected_runs[-1]
    
    prefix = paragraph.runs[first_run_idx].text[:start_idx - run_offsets[first_run_idx][0]]
    suffix = paragraph.runs[last_run_idx].text[end_idx - run_offsets[last_run_idx][0]:]
    
    paragraph.runs[first_run_idx].text = prefix + new_text
    if track:
        paragraph.runs[first_run_idx].font.highlight_color = WD_COLOR_INDEX.YELLOW
        
    for idx in range(first_run_idx + 1, last_run_idx):
        paragraph.runs[idx].text = ""
        
    if last_run_idx != first_run_idx:
        paragraph.runs[last_run_idx].text = suffix
        
    return True

def replace_hypothesis(paragraph, k, new_text, track=False):
    full_text = "".join(r.text for r in paragraph.runs).strip()
    if full_text.startswith(k) or full_text.startswith(k.replace('۱', '1')) or full_text.startswith(k.replace('۲', '2')) or full_text.startswith(k.replace('۳', '3')) or full_text.startswith(k.replace('۴', '4')) or full_text.startswith(k.replace('۵', '5')):
        # Replace the entire text of the paragraph by replacing the first run and clearing others
        paragraph.runs[0].text = new_text
        if track:
            paragraph.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
        for idx in range(1, len(paragraph.runs)):
            paragraph.runs[idx].text = ""
        return True
    return False

def process_document(input_path, output_clean, output_tracked):
    doc_clean = docx.Document(input_path)
    doc_tracked = docx.Document(input_path)
    
    old_c3_1 = "و تاثیر آن بر همدلی"
    old_c3_2 = "و تأثیر آن بر همدلی"
    new_c3 = "و نقش آن در ارتقای همدلی"
    
    hypotheses = {
        "۱.": "۱. بین سبکهای دلبستگی و همدلی در دانشجویان پزشکی دانشگاه علوم پزشکی قم رابطه معناداری وجود دارد (سبکهای دلبستگی پیشبینیکننده معنادار همدلی است).",
        "۲.": "۲. بین سبکهای دلبستگی و تنظیم هیجان در دانشجویان پزشکی دانشگاه علوم پزشکی قم رابطه معناداری وجود دارد (سبکهای دلبستگی پیشبینیکننده معنادار تنظیم هیجان است).",
        "۳.": "۳. بین سبکهای دلبستگی و خودشفقتی در دانشجویان پزشکی دانشگاه علوم پزشکی قم رابطه معناداری وجود دارد (سبکهای دلبستگی پیشبینیکننده معنادار خودشفقتی است).",
        "۴.": "۴. بین تنظیم هیجان و همدلی در دانشجویان پزشکی دانشگاه علوم پزشکی قم رابطه معناداری وجود دارد (تنظیم هیجان پیشبینیکننده معنادار همدلی است).",
        "۵.": "۵. بین خودشفقتی و همدلی در دانشجویان پزشکی دانشگاه علوم پزشکی قم رابطه معناداری وجود دارد (خودشفقتی پیشبینیکننده معنادار همدلی است)."
    }
    
    def apply_changes(doc, track):
        in_hypotheses = False
        for para in doc.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            
            replace_in_runs(para, old_c3_1, new_c3, track)
            replace_in_runs(para, old_c3_2, new_c3, track)
            replace_in_runs(para, "تاثیر آن بر همدلی", "نقش آن در ارتقای همدلی", track)
            replace_in_runs(para, "تأثیر آن بر همدلی", "نقش آن در ارتقای همدلی", track)
            
            if "فرضیات" in text or "سؤالات پژوهش" in text:
                in_hypotheses = True
                
            if in_hypotheses:
                for k, v in hypotheses.items():
                    if text.startswith(k) or text.startswith(k.replace('۱', '1').replace('۲', '2').replace('۳', '3').replace('۴', '4').replace('۵', '5')):
                        replace_hypothesis(para, k, v, track)
                        
    apply_changes(doc_clean, track=False)
    apply_changes(doc_tracked, track=True)
    
    doc_clean.save(output_clean)
    doc_tracked.save(output_tracked)

if __name__ == "__main__":
    input_file = "/home/saber-ghaderi/My Work/Mehrane/01_raw_inputs/فایل پروپوزال.docx"
    output_c = "/home/saber-ghaderi/My Work/Mehrane/03_deliverables/Proposal_Revised_Clean.docx"
    output_t = "/home/saber-ghaderi/My Work/Mehrane/03_deliverables/Proposal_Revised_Tracked.docx"
    
    process_document(input_file, output_c, output_t)
