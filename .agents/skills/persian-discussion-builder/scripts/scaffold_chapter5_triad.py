#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AcademicSuite Chapter 5 Triad Artifact Scaffolder & Compiler (scaffold_chapter5_triad.py)
------------------------------------------------------------------------------------------
Compiles the synchronized triad (.docx, .md, .json) for any Chapter 5 micro-stage
strictly using content PRODUCED BY THE AGENT.

CONSTITUTIONAL MANDATE (Zero Template / Zero Prewritten Text):
- Deterministic scripts ('The Hands') MUST NOT contain hardcoded prewritten paragraphs
  or canned boilerplate essay templates.
- Scholarly text formulation belongs exclusively to autonomous cognitive subagents
  ('academic-writer', 'literature-expert').
- This script compiles the agent's authored markdown narrative into institutional
  OpenXML .docx (B Titr / B Nazanin typography, RTL <w:bidi>, decoupled LTR stats)
  and outputs the synchronized verification triad.
"""

import os
import sys
import json
import hashlib
import argparse
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, Optional

# Dynamic discovery of local virtualenv site-packages (.venv / venv)
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../.."))
for venv_name in [".venv", "venv"]:
    venv_lib = os.path.join(ROOT_DIR, venv_name, "lib")
    if os.path.isdir(venv_lib):
        for entry in os.listdir(venv_lib):
            sp = os.path.join(venv_lib, entry, "site-packages")
            if os.path.isdir(sp) and sp not in sys.path:
                sys.path.insert(0, sp)

for p in ["/usr/lib/python3/dist-packages", "/usr/local/lib/python3/dist-packages"]:
    if os.path.exists(p) and p not in sys.path:
        sys.path.append(p)

try:
    import docx
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def compute_sha256(filepath: str) -> str:
    """Computes SHA-256 hash of a file."""
    if not os.path.isfile(filepath):
        return ""
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def clean_whitespace(text: str) -> str:
    """Normalize whitespace without altering paragraph structure."""
    if not text:
        return ""
    return text.strip()


def set_bidi_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Enforce Persian BiDi RTL directionality and alignment."""
    pPr = p._p.get_or_add_pPr()
    if not pPr.xpath('./w:bidi'):
        bidi = parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>')
        pPr.insert(0, bidi)
    if align == WD_ALIGN_PARAGRAPH.RIGHT:
        jc = pPr.find(qn('w:jc'))
        if jc is not None:
            pPr.remove(jc)
    else:
        p.alignment = align


def add_text_run(p, text, font_fa='B Nazanin', font_en='Times New Roman', size=13, bold=False, italic=False):
    """Add text run with explicit Persian/Latin font bindings and complex script tags."""
    run = p.add_run(str(text))
    run.font.name = font_fa
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

    rPr = run._r.get_or_add_rPr()
    sz_val = int(size * 2)
    has_persian = any('\u0600' <= ch <= '\u06FF' or '\uFB50' <= ch <= '\uFDFF' or '\uFE70' <= ch <= '\uFEFF' for ch in str(text))
    
    if has_persian:
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} '
            f'w:ascii="{font_fa}" w:hAnsi="{font_fa}" '
            f'w:cs="{font_fa}" w:eastAsia="{font_fa}" w:hint="cs"/>'
        )
        rPr.append(rFonts)
        rtl = parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>')
        rPr.append(rtl)
    else:
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} '
            f'w:ascii="{font_en}" w:hAnsi="{font_en}" '
            f'w:cs="{font_fa}" w:eastAsia="{font_fa}"/>'
        )
        rPr.append(rFonts)

    szCs = parse_xml(f'<w:szCs {nsdecls("w")} w:val="{sz_val}"/>')
    rPr.append(szCs)
    if bold:
        bCs = parse_xml(f'<w:bCs {nsdecls("w")} w:val="1"/>')
        rPr.append(bCs)
    return run


def compile_chapter5_triad(
    stage_id: str,
    base_name: str,
    output_dir: str,
    stage_title: str,
    agent_narrative: Optional[str] = None,
    narrative_file: Optional[str] = None,
    articles_corpus_path: Optional[str] = None,
    metadata_payload: Optional[Dict[str, Any]] = None
) -> Dict[str, str]:
    """
    Compiles the synchronized triad (.docx, .md, .json) using narrative PRODUCED BY THE AGENT.
    Zero prewritten text templates allowed.
    """
    out_dir_path = Path(output_dir)
    out_dir_path.mkdir(parents=True, exist_ok=True)

    # 1. Resolve agent narrative text
    narrative_text = ""
    if agent_narrative:
        narrative_text = agent_narrative.strip()
    elif narrative_file and os.path.exists(narrative_file):
        with open(narrative_file, "r", encoding="utf-8") as f:
            narrative_text = f.read().strip()
    else:
        # Check if draft.md already exists in output_dir
        existing_draft = out_dir_path / f"{base_name}.md"
        if existing_draft.exists():
            with open(existing_draft, "r", encoding="utf-8") as f:
                narrative_text = f.read().strip()

    if not narrative_text or len(narrative_text.strip()) < 10:
        raise ValueError(
            f"Zero-Template Invariant Violation (Directives 0 & 3): Agent-authored narrative is required for stage '{stage_id}'. "
            "AcademicSuite prohibits static template paragraphs and canned boilerplate. Provide narrative via --narrative or --narrative-file."
        )

    # Ensure title heading exists at the top
    if not narrative_text.startswith("# "):
        full_md_content = f"# {stage_title}\n\n{narrative_text}\n"
    else:
        full_md_content = narrative_text + "\n"
    status = "STAGE_DRAFTED"

    # Ingest reference citations if an articles corpus was provided
    referenced_articles = []
    if articles_corpus_path and os.path.exists(articles_corpus_path):
        try:
            with open(articles_corpus_path, "r", encoding="utf-8") as f:
                cdata = json.load(f)
            # Scan agent narrative for mentioned authors/citations
            for art in cdata.get("articles", []):
                author = art.get("author", "")
                year = str(art.get("year", ""))
                if author and (author.lower() in full_md_content.lower() or year in full_md_content):
                    referenced_articles.append({
                        "citation": art.get("apa_citation"),
                        "bib": art.get("apa_bib"),
                        "doi": art.get("doi")
                    })
        except Exception:
            pass

    # 2. Triad JSON Deliverable
    json_path = out_dir_path / f"{base_name}.json"
    json_data = {
        "stage_id": stage_id,
        "base_name": base_name,
        "stage_title": stage_title,
        "compiled_at": datetime.now(timezone.utc).isoformat(),
        "status": status,
        "char_count": len(narrative_text),
        "word_count": len(narrative_text.split()),
        "metrics": {
            "char_count": len(narrative_text),
            "word_count": len(narrative_text.split())
        },
        "referenced_articles": referenced_articles,
        "parameters": metadata_payload or {}
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_data, f, indent=2, ensure_ascii=False)

    # 3. Triad Markdown Deliverable
    md_path = out_dir_path / f"{base_name}.md"
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(full_md_content)

    # 4. Triad OpenXML DOCX Deliverable
    docx_path = out_dir_path / f"{base_name}.docx"
    if HAS_DOCX:
        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.right_margin = Inches(1.18)
            section.left_margin = Inches(1.0)

        lines = full_md_content.split("\n")
        for line in lines:
            line_str = clean_whitespace(line)
            if not line_str or line_str.startswith("<!--"):
                continue
            p = doc.add_paragraph()
            if line_str.startswith("# "):
                set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
                add_text_run(p, line_str.replace("# ", ""), font_fa="B Titr", size=14, bold=True)
            elif line_str.startswith("## "):
                set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
                add_text_run(p, line_str.replace("## ", ""), font_fa="B Titr", size=13, bold=True)
            elif line_str.startswith("### "):
                set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
                add_text_run(p, line_str.replace("### ", ""), font_fa="B Titr", size=12, bold=True)
            elif line_str.startswith("**") and line_str.endswith("**"):
                set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
                add_text_run(p, line_str.replace("**", ""), font_fa="B Nazanin", size=13, bold=True)
            else:
                set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
                add_text_run(p, line_str, font_fa="B Nazanin", size=13)

        doc.save(str(docx_path))
    else:
        with open(docx_path, "wb") as f:
            f.write(b"MOCK_DOCX_STAGE_TRIAD")

    # Update JSON deliverable with deliverables manifest and SHA256 hashes
    json_data["deliverables"] = {
        "json": str(json_path.resolve()),
        "json_sha256": compute_sha256(str(json_path)),
        "md": str(md_path.resolve()),
        "md_sha256": compute_sha256(str(md_path)),
        "docx": str(docx_path.resolve()),
        "docx_sha256": compute_sha256(str(docx_path)),
        "sha256": compute_sha256(str(json_path))
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_data, f, indent=2, ensure_ascii=False)

    print(f"Compiled Chapter 5 Triad [{stage_id}]:")
    print(f"  - Status: {status}")
    print(f"  - JSON:   {json_path}")
    print(f"  - MD:     {md_path}")
    print(f"  - DOCX:   {docx_path}")

    return {
        "stage_id": stage_id,
        "json": str(json_path.resolve()),
        "md": str(md_path.resolve()),
        "docx": str(docx_path.resolve()),
        "status": status
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Compile Chapter 5 Stage Triad from Agent Narrative")
    parser.add_argument("--stage", required=True, help="Stage ID (e.g. 01_findings_recap, 02_hypothesis_1_discussion)")
    parser.add_argument("--base", required=True, help="Base filename (e.g. 01_findings_recap)")
    parser.add_argument("--outdir", default=".", help="Output directory path")
    parser.add_argument("--title", required=True, help="Stage title in Persian")
    parser.add_argument("--narrative", default=None, help="Agent-authored narrative string")
    parser.add_argument("--narrative-file", default=None, help="Path to markdown file containing agent-authored narrative")
    parser.add_argument("--articles", default=None, help="Path to article_enrichment_cards.json")
    args = parser.parse_args()

    compile_chapter5_triad(
        stage_id=args.stage,
        base_name=args.base,
        output_dir=args.outdir,
        stage_title=args.title,
        agent_narrative=args.narrative,
        narrative_file=args.narrative_file,
        articles_corpus_path=args.articles
    )
