#!/usr/bin/env python3
"""
Persian Chapter 5 Word Document Generator (generate_chapter5_docx.py)
---------------------------------------------------------------------
Generates standard Iranian graduate school Chapter 5 (بحث و نتیجه‌گیری) Word documents (.docx)
with correct OpenXML RTL directionality and Persian typography (B Nazanin, B Titr).

Sections:
1. مقدمه (Introduction)
2. بحث و بررسی فرضیه‌ها (Hypothesis-by-Hypothesis Discussion)
3. پیامدهای کاربردی و بالینی (Clinical & Applied Implications)
4. محدودیت‌های پژوهش (Limitations)
5. پیشنهادهای پژوهش (Research & Practical Recommendations)
6. نتیجه‌گیری نهایی (Final Conclusion)
"""

import os
import sys
import json
import argparse
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_paragraph_bidi(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Enforce Persian BiDi RTL directionality and alignment on paragraph."""
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    if not pPr.xpath('./w:bidi'):
        bidi = parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>')
        pPr.insert(0, bidi)

def add_run(p, text, font_fa='B Nazanin', font_en='Times New Roman', size=13, bold=False, italic=False):
    """Add text run with explicit Persian/Latin font bindings, w:rtl, and complex script formatting."""
    run = p.add_run(str(text))
    run.font.name = font_fa
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    
    rPr = run._r.get_or_add_rPr()
    sz_val = int(size * 2)
    has_persian = any('\u0600' <= ch <= '\u06FF' or '\uFB50' <= ch <= '\uFDFF' or '\uFE70' <= ch <= '\uFEFF' for ch in str(text))
    if has_persian:
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} '
            f'w:ascii="{font_fa}" w:hAnsi="{font_fa}" '
            f'w:cs="{font_fa}" w:eastAsia="{font_fa}" w:hint="cs"/>'
        )
        rPr.append(rFonts)
        rtl = parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>')
        rPr.append(rtl)
    else:
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} '
            f'w:ascii="{font_en}" w:hAnsi="{font_en}" '
            f'w:cs="{font_fa}" w:eastAsia="{font_fa}"/>'
        )
        rPr.append(rFonts)

    szCs = parse_xml(f'<w:szCs {nsdecls("w")} w:val="{sz_val}"/>')
    rPr.append(szCs)
    if bold:
        bCs = parse_xml(f'<w:bCs {nsdecls("w")} w:val="1"/>')
        rPr.append(bCs)
    return run

def build_chapter5_document(data: dict, output_path: str):
    doc = docx.Document()
    
    # Page Margins (3 cm right, 2.5 cm others)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.right_margin = Inches(1.18)
        section.left_margin = Inches(1.0)
        
    # --- Chapter Title ---
    p_ch = doc.add_paragraph()
    set_paragraph_bidi(p_ch, WD_ALIGN_PARAGRAPH.CENTER)
    p_ch.paragraph_format.space_before = Pt(24)
    p_ch.paragraph_format.space_after = Pt(12)
    add_run(p_ch, "فصل پنجم", font_fa='B Titr', size=18, bold=True)
    
    p_title = doc.add_paragraph()
    set_paragraph_bidi(p_title, WD_ALIGN_PARAGRAPH.CENTER)
    p_title.paragraph_format.space_after = Pt(28)
    add_run(p_title, "بحث و نتیجه‌گیری", font_fa='B Titr', size=16, bold=True)
    
    # --- 1-5. Introduction ---
    p_h1 = doc.add_paragraph()
    set_paragraph_bidi(p_h1, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)
    add_run(p_h1, "۱-۵. مقدمه", font_fa='B Titr', size=14, bold=True)
    
    intro_text = data.get("introduction", (
        "پژوهش حاضر با هدف بررسی روابط ساختاری و ارزیابی اثربخشی متغیرهای پژوهش در جامعه هدف به مرحله اجرا درآمد. "
        "پس از جمع‌آوری داده‌ها، سنجش فرضیه‌ها و تجزیه‌وتحلیل آماری در فصل چهارم، در این فصل یافته‌های پژوهش مورد "
        "بحث و تفسیر جامع قرار می‌گیرد. ساختار فصل حاضر به‌گونه‌ای تدوین شده است که ابتدا یافته‌های مربوط به هریک "
        "از فرضیه‌های پژوهش به‌صورت مجزا با پژوهش‌های پیشین همسو و ناهمسو مقایسه گردیده و سپس بر مبنای سازوکارهای "
        "نظری و روان‌شناختی تبیین می‌گردد. در ادامه، پیامدهای بالینی و کاربردی، محدودیت‌های روش‌شناختی پژوهش و در نهایت "
        "پیشنهادهای پژوهشی و کاربردی ارائه خواهد شد."
    ))
    p_intro = doc.add_paragraph()
    set_paragraph_bidi(p_intro, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_intro.paragraph_format.line_spacing = 1.25
    add_run(p_intro, intro_text)
    
    # --- 2-5. Hypothesis Discussions ---
    p_h2 = doc.add_paragraph()
    set_paragraph_bidi(p_h2, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h2.paragraph_format.space_before = Pt(16)
    p_h2.paragraph_format.space_after = Pt(6)
    add_run(p_h2, "۲-۵. بحث و بررسی پیرامون یافته‌های حاصل از فرضیه‌ها", font_fa='B Titr', size=14, bold=True)
    
    hypotheses = data.get("hypotheses", [])
    for idx, hyp in enumerate(hypotheses):
        p_hyp_title = doc.add_paragraph()
        set_paragraph_bidi(p_hyp_title, WD_ALIGN_PARAGRAPH.RIGHT)
        p_hyp_title.paragraph_format.space_before = Pt(12)
        p_hyp_title.paragraph_format.space_after = Pt(4)
        add_run(p_hyp_title, f"۱-۲-۵. بررسی و تبیین {hyp.get('title', f'فرضیه شماره {idx+1}')}", font_fa='B Nazanin', size=13, bold=True)
        
        # Narrative components: statement, empirical alignment, theoretical mechanism
        full_text = (
            f"{hyp.get('statement', '')} "
            f"{hyp.get('empirical_comparison', '')} "
            f"{hyp.get('theoretical_explanation', '')}"
        )
        p_hyp_body = doc.add_paragraph()
        set_paragraph_bidi(p_hyp_body, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p_hyp_body.paragraph_format.line_spacing = 1.25
        add_run(p_hyp_body, full_text.strip())
        
    # --- 3-5. Implications ---
    p_h3 = doc.add_paragraph()
    set_paragraph_bidi(p_h3, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h3.paragraph_format.space_before = Pt(16)
    p_h3.paragraph_format.space_after = Pt(6)
    add_run(p_h3, "۳-۵. پیامدهای کاربردی و بالینی پژوهش", font_fa='B Titr', size=14, bold=True)
    
    implications = data.get("implications", (
        "یافته‌های این پژوهش واجد دلالت‌های نظری و کاربردی متعددی در حوزه‌های مشاوره، روان‌درمانی و برنامه‌ریزی‌های آموزشی است. "
        "در حوزه بالینی، شناسایی متغیرهای تعدیل‌کننده و میانجی به متخصصان سلامت روان این امکان را می‌دهد تا بسته‌های درمانی "
        "و مشاوره‌ای متناسب‌تری را تدوین نمایند. همچنین برای مراکز مشاوره، مؤسسات آموزشی و دانشگاه‌ها توصیه می‌شود با برگزاری "
        "کارگاه‌های ارتقای مهارت‌های شناختی و هیجانی، بستر مناسبی را جهت کاهش آسیب‌پذیری روانی افراد فراهم آورند."
    ))
    p_imp = doc.add_paragraph()
    set_paragraph_bidi(p_imp, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_imp.paragraph_format.line_spacing = 1.25
    add_run(p_imp, implications)
    
    # --- 4-5. Limitations ---
    p_h4 = doc.add_paragraph()
    set_paragraph_bidi(p_h4, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h4.paragraph_format.space_before = Pt(16)
    p_h4.paragraph_format.space_after = Pt(6)
    add_run(p_h4, "۴-۵. محدودیت‌های پژوهش", font_fa='B Titr', size=14, bold=True)
    
    limitations = data.get("limitations", [
        "استفاده از روش نمونه‌گیری در دسترس که تعمیم‌پذیری یافته‌ها به سایر جوامع و گروه‌های جمعیتی را با احتیاط مواجه می‌سازد.",
        "اتکای صرف به ابزارهای خودگزارش‌دهی که همواره احتمال سوگیری مطلوبیت اجتماعی و عدم دقت کامل آزمودنی‌ها را به همراه دارد.",
        "عدم امکان اجرای دوره پیگیری (۳ یا ۶ ماهه) جهت سنجش پایداری اثرات مداخله در طول زمان به دلیل محدودیت‌های زمانی.",
        "عدم کنترل کامل متغیرهای ناخواسته نظیر تفاوت‌های فردی در سبک‌های والدگری و سطح اقتصادی-اجتماعی آزمودنی‌ها."
    ])
    for lim in limitations:
        p_lim = doc.add_paragraph()
        set_paragraph_bidi(p_lim, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p_lim.paragraph_format.line_spacing = 1.25
        p_lim.paragraph_format.left_indent = Inches(0.2)
        add_run(p_lim, f"• {lim}")
        
    # --- 5-5. Recommendations ---
    p_h5 = doc.add_paragraph()
    set_paragraph_bidi(p_h5, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h5.paragraph_format.space_before = Pt(16)
    p_h5.paragraph_format.space_after = Pt(6)
    add_run(p_h5, "۵-۵. پیشنهادهای پژوهش", font_fa='B Titr', size=14, bold=True)
    
    # 5-5-1 Research Recommendations
    p_h5_1 = doc.add_paragraph()
    set_paragraph_bidi(p_h5_1, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h5_1.paragraph_format.space_before = Pt(10)
    p_h5_1.paragraph_format.space_after = Pt(4)
    add_run(p_h5_1, "۱-۵-۵. پیشنهادهای پژوهشی (برای محققان آینده)", font_fa='B Nazanin', size=13, bold=True)
    
    res_recs = data.get("research_recommendations", [
        "پیشنهاد می‌شود در پژوهش‌های آتی از روش‌های نمونه‌گیری تصادفی خوشه‌ای یا چندمرحله‌ای جهت افزایش توان تعمیم‌پذیری استفاده گردد.",
        "انجام پژوهش‌های مشابه با به‌کارگیری ابزارهای سنجش چندگانه (نظیر مصاحبه‌های بالینی ساختاریافته و شاخص‌های فیزیولوژیک) توصیه می‌شود.",
        "پیشنهاد می‌گردد مدل حاضر در جوامع هدف دیگر نظیر گروه‌های سنی مختلف و با در نظر گرفتن متغیر جنسیت به‌عنوان تعدیل‌کننده مورد بررسی قرار گیرد."
    ])
    for r in res_recs:
        p_r = doc.add_paragraph()
        set_paragraph_bidi(p_r, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p_r.paragraph_format.line_spacing = 1.25
        p_r.paragraph_format.left_indent = Inches(0.2)
        add_run(p_r, f"• {r}")
        
    # 5-5-2 Practical Recommendations
    p_h5_2 = doc.add_paragraph()
    set_paragraph_bidi(p_h5_2, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h5_2.paragraph_format.space_before = Pt(10)
    p_h5_2.paragraph_format.space_after = Pt(4)
    add_run(p_h5_2, "۲-۵-۵. پیشنهادهای کاربردی (برای سازمان‌ها و درمانگران)", font_fa='B Nazanin', size=13, bold=True)
    
    prac_recs = data.get("practical_recommendations", [
        "به مشاوران و روان‌درمانگران توصیه می‌شود در فرایند مداخله درمانی، بر تقویت سازه‌های محافظت‌کننده شناسایی‌شده در این پژوهش تمرکز نمایند.",
        "پیشنهاد می‌گردد مراکز مشاوره دانشجویی دوره‌های آموزشی مبتنی بر یافته‌های این پژوهش را جهت توانمندسازی مراجعان طراحی و برگزار کنند."
    ])
    for pr in prac_recs:
        p_pr = doc.add_paragraph()
        set_paragraph_bidi(p_pr, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p_pr.paragraph_format.line_spacing = 1.25
        p_pr.paragraph_format.left_indent = Inches(0.2)
        add_run(p_pr, f"• {pr}")
        
    # --- 6-5. Final Conclusion ---
    p_h6 = doc.add_paragraph()
    set_paragraph_bidi(p_h6, WD_ALIGN_PARAGRAPH.RIGHT)
    p_h6.paragraph_format.space_before = Pt(16)
    p_h6.paragraph_format.space_after = Pt(6)
    add_run(p_h6, "۶-۵. نتیجه‌گیری نهایی", font_fa='B Titr', size=14, bold=True)
    
    conclusion = data.get("conclusion", (
        "در نهایت می‌توان نتیجه گرفت که متغیرهای مورد بررسی در پژوهش حاضر نقش تعیین‌کننده و معناداری در ارتقای بهزیستی روان‌شناختی "
        "و تعدیل نشانه‌های آسیب‌شناختی ایفا می‌کنند. یافته‌های این رساله نه‌تنها شواهد تجربی نیرومندی در تأیید الگوهای نظری موجود "
        "فراهم آورد، بلکه پنجره‌های نوینی را فراروی مداخلات درمانی هدفمند در حوزه سلامت روان گشود."
    ))
    p_conc = doc.add_paragraph()
    set_paragraph_bidi(p_conc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_conc.paragraph_format.line_spacing = 1.25
    add_run(p_conc, conclusion)
    
    doc.save(output_path)
    print(f"Chapter 5 Document successfully generated at: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Chapter 5 Persian Word Document Generator")
    parser.add_argument("--json", required=True, help="Path to JSON file containing Chapter 5 discussion content")
    parser.add_argument("--out", default="Chapter_5_Discussion_and_Conclusion.docx", help="Output .docx file path")
    args = parser.parse_args()
    
    with open(args.json, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    build_chapter5_document(data, args.out)

if __name__ == "__main__":
    main()
