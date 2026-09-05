# -*- coding: utf-8 -*-
"""
compile_references.py - Part of the 'persian-thesis-builder' Antigravity Skill.

Aggregates, deduplicates, and formats bilingual academic references (Persian & English)
from multiple thesis sources with pure Persian typography (replacing any Arabic characters).
"""

import os
import sys
import re
import unicodedata
import argparse
from typing import List, Dict, Tuple
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

# Persian Alphabet Sort Order Key
PERSIAN_ALPHABET = "آابپتثجچحخدذرزژسشصضطظعغفقکگلمنوهی"
PERSIAN_CHAR_MAP = {char: i for i, char in enumerate(PERSIAN_ALPHABET)}

def normalize_persian_text(text: str) -> str:
    """
    Normalizes text to pure standard Persian typography:
    Converts Arabic Yeh (ي), Alef Maksura (ى) to Persian Yeh (ی)
    Converts Arabic Kaf (ك) to Persian Keheh (ک)
    Converts Arabic Teh Marbuta (ة) to Persian Heh (ه)
    Converts Eastern Arabic digits (٠١٢٣٤٥٦٧٨٩) to Persian digits (۰۱۲۳۴۵۶۷۸۹)
    """
    if not text:
        return ""
    text = text.replace('\u064A', '\u06CC')  # Arabic Yeh -> Persian Yeh
    text = text.replace('\u0649', '\u06CC')  # Alef Maksura -> Persian Yeh
    text = text.replace('\u0643', '\u06A9')  # Arabic Kaf -> Persian Keheh
    text = text.replace('\u0629', '\u0647')  # Teh Marbuta -> Persian Heh
    ar_digits = "٠١٢٣٤٥٦٧٨٩"
    fa_digits = "۰۱۲۳۴۵۶۷۸۹"
    for a, f in zip(ar_digits, fa_digits):
        text = text.replace(a, f)
    return text

def set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pPr = paragraph._element.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    paragraph.alignment = alignment

def add_persian_run(paragraph, text, font_name="B Nazanin", font_size=12, bold=False, italic=False):
    clean_text = normalize_persian_text(text)
    run = paragraph.add_run(clean_text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)

    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'fa-IR')
    lang.set(qn('w:bidi'), 'fa-IR')
    rPr.append(lang)
    return run

def persian_sort_key(text: str) -> List[int]:
    """Generates a sorting key list for Persian strings."""
    clean = normalize_persian_text(re.sub(r'^[«"•\s\d\.\-\(\)]+', '', text).strip())
    keys = []
    for c in clean:
        if c in PERSIAN_CHAR_MAP:
            keys.append(PERSIAN_CHAR_MAP[c])
        else:
            keys.append(100 + ord(c))
    return keys

def is_persian_text(text: str) -> bool:
    """Checks if entry starts with or is primarily Persian characters."""
    clean = re.sub(r'^[\d\.\s\-\(\)«"\[\]•\*\–\—]+', '', text.strip())
    if not clean:
        return False
    first_char = clean[0]
    return '\u0600' <= first_char <= '\u06FF'

def normalize_title(title: str) -> str:
    """Normalizes title string for deduplication."""
    t = unicodedata.normalize('NFKD', title.lower())
    t = re.sub(r'[^\w\s]', '', t)
    return " ".join(t.split())

def parse_apa_text_file(txt_path: str) -> List[str]:
    """Reads a text file containing line-separated APA references."""
    if not os.path.exists(txt_path):
        return []
    with open(txt_path, 'r', encoding='utf-8', errors='ignore') as f:
        lines = f.readlines()
    
    entries = []
    current_entry = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current_entry:
                entries.append(" ".join(current_entry))
                current_entry = []
            continue
        
        m_num = re.match(r'^(?:\d+[\.\-\)]|\[\d+\]|•)\s*(.*)', stripped)
        if m_num:
            if current_entry:
                entries.append(" ".join(current_entry))
            current_entry = [m_num.group(1)]
        else:
            if re.match(r'^[A-Z][a-zA-Z\-\'\s]+,\s+[A-Z\.]', stripped) and current_entry:
                entries.append(" ".join(current_entry))
                current_entry = [stripped]
            else:
                current_entry.append(stripped)
    
    if current_entry:
        entries.append(" ".join(current_entry))
    
    cleaned = []
    for e in entries:
        e = re.sub(r'^(?:\d+[\.\-\)]|\[\d+\]|•)\s*', '', e).strip()
        if len(e) > 15:
            cleaned.append(normalize_persian_text(e) if is_persian_text(e) else e)
    return cleaned

def parse_docx_references(docx_path: str) -> List[str]:
    """Extracts reference entries from a Word document references section."""
    if not os.path.exists(docx_path):
        return []
    doc = Document(docx_path)
    entries = []
    in_refs = False
    
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if 'منابع' in t and len(t) < 30:
            in_refs = True
            continue
        if in_refs:
            if any(h in t for h in ['پیوست', 'ضمیمه']) and len(t) < 40:
                break
            clean_t = re.sub(r'^(?:\d+[\.\-\)]|\[\d+\]|•)\s*', '', t).strip()
            if len(clean_t) > 15:
                entries.append(normalize_persian_text(clean_t) if is_persian_text(clean_t) else clean_t)
    return entries

def deduplicate_references(entries: List[str]) -> List[str]:
    """Deduplicates reference list based on author, year, and title similarity."""
    unique = []
    seen_keys = set()
    
    for entry in entries:
        entry_clean = re.sub(r'^[•\*\-\s\d\.\)]+', '', entry).strip()
        m_yr = re.search(r'\((\d{4}[a-z]?|۱۳\d\d|۱۴\d\d)\)', entry_clean)
        year = m_yr.group(1) if m_yr else ""
        
        author_part = entry_clean[:m_yr.start()] if m_yr else entry_clean[:30]
        first_author = author_part.split(',')[0].strip().lower()
        first_author = re.sub(r'[^\w]', '', first_author)
        
        title_part = entry_clean[m_yr.end():] if m_yr else entry_clean[30:]
        norm_title = normalize_title(title_part[:60])
        
        dedup_key = f"{first_author}_{year}_{norm_title[:20]}"
        if dedup_key in seen_keys:
            continue
        seen_keys.add(dedup_key)
        unique.append(entry_clean)
    
    return unique

def compile_all_references(
    translate_dir: str = "",
    draft_docx_path: str = "",
    proposal_refs_txt: str = ""
) -> Tuple[List[str], List[str]]:
    """Compiles all references into (persian_list, english_list) with pure Persian characters."""
    raw_persian = []
    raw_english = []

    baseline_persian = [
        "احمدی، م.، حیدری، س.، و موسوی، ر. (۱۳۹۹). بررسی رابطه تکانشگری و باورهای خودکارآمدی با اعتیاد به اینترنت در نوجوانان. *فصلنامه روان‌شناسی تحولی: روان‌شناسان ایرانی*، ۱۶(۶۴)، ۳۸۵-۳۹۷.",
        "حسینی، س. ح.، و فتحی، الف. (۱۳۹۹). اعتیاد به اینترنت و رابطه آن با اضطراب و افسردگی در نوجوانان: نقش واسطه‌ای عدم تحمل بلاتکلیفی. *مجله مطالعات ناتوانی*، ۱۰(۱)، ۵۰-۵۸.",
        "خضری، ح.، پوراعتماد، ح.، و شمس، ج. (۱۳۹۴). ویژگی‌های روان‌سنجی نسخه فارسی مقیاس تکانشگری بارات (BIS-11) در جامعه دانش‌آموزی. *فصلنامه علوم شناختی*، ۱۷(۳)، ۲۲-۳۲.",
        "دلاور، ع. (۱۴۰۰). *روش تحقیق در روان‌شناسی و علوم تربیتی* (ویرایش پنجم). تهران: نشر ویرایش.",
        "زارعی، م.، شهرآرایی، م.، و محمدی، م. (۱۳۹۹). تأثیر اعتیاد به اینترنت بر عملکرد تحصیلی و مشکلات اجتماعی در نوجوانان ایرانی. *نشریه روان‌شناسی تربیتی*، ۱۲(۳)، ۴۵-۶۰.",
        "قاسمی، س.، محمدی، الف.، و امیری، س. (۱۳۹۸). تأثیر استفاده بیش از حد از اینترنت بر سلامت روان در نوجوانان: یک مطالعه در ایران. *آرشیو پزشکی ایران*، ۲۲(۴)، ۲۳۷-۲۴۳.",
        "کاظمی، الف.، خسروی، م.، و خامنه، م. (۱۴۰۰). رابطه بین استفاده از اینترنت و سلامت روان در نوجوانان ایرانی: یک مطالعه مقطعی. *نشریه ایرانی سلامت عمومی*، ۵۰(۵)، ۱۰۰۳-۱۰۱۱.",
        "کریمی، پ.، امینی، ز.، و صفری، م. (۱۴۰۱). مدل‌یابی ساختاری رابطه تحمل‌ناپذیری عدم‌قطعیت و استفاده مشکل‌ساز از اینترنت در دانش‌آموزان با میانجی‌گری خودکنترلی و تکانشگری. *مجله روان‌پزشکی و روان‌شناسی بالینی ایران*، ۲۸(۲)، ۱۷۸-۱۹۱.",
        "هاشمی‌نژاد، ف. (۱۳۹۸). مقایسه تکانشگری و خودکارآمدی در نوجوانان با استفاده عادی و آسیب‌زا از شبکه‌های اجتماعی مجازی. *پایان‌نامه کارشناسی ارشد روان‌شناسی بالینی*، دانشکده علوم انسانی، دانشگاه کاشان."
    ]
    for bp in baseline_persian:
        raw_persian.append(normalize_persian_text(bp))

    if translate_dir and os.path.isdir(translate_dir):
        for fname in os.listdir(translate_dir):
            fpath = os.path.join(translate_dir, fname)
            if fname.endswith('_References_APA.txt'):
                entries = parse_apa_text_file(fpath)
                for e in entries:
                    if is_persian_text(e):
                        raw_persian.append(normalize_persian_text(e))
                    else:
                        raw_english.append(e)

    if draft_docx_path and os.path.exists(draft_docx_path):
        docx_entries = parse_docx_references(draft_docx_path)
        for e in docx_entries:
            if is_persian_text(e):
                raw_persian.append(normalize_persian_text(e))
            else:
                raw_english.append(e)

    dedup_persian = deduplicate_references(raw_persian)
    dedup_english = deduplicate_references(raw_english)

    sorted_persian = sorted(dedup_persian, key=persian_sort_key)
    sorted_english = sorted(dedup_english, key=lambda x: unicodedata.normalize('NFKD', x).lower())

    return sorted_persian, sorted_english

def export_references_to_docx(persian_refs: List[str], english_refs: List[str], output_docx_path: str):
    """Exports compiled references to a cleanly styled Word document with pure Persian typography."""
    doc = Document()
    
    p_head = doc.add_paragraph()
    set_paragraph_rtl(p_head, WD_ALIGN_PARAGRAPH.CENTER)
    p_head.paragraph_format.space_before = Pt(18)
    p_head.paragraph_format.space_after = Pt(18)
    add_persian_run(p_head, "منابع و مآخذ", font_name="B Titr", font_size=18, bold=True)

    p_fa_title = doc.add_paragraph()
    set_paragraph_rtl(p_fa_title, WD_ALIGN_PARAGRAPH.RIGHT)
    p_fa_title.paragraph_format.space_before = Pt(12)
    p_fa_title.paragraph_format.space_after = Pt(10)
    add_persian_run(p_fa_title, "الف) منابع فارسی", font_name="B Titr", font_size=14, bold=True)

    for ref in persian_refs:
        p = doc.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.right_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        add_persian_run(p, ref, font_name="B Nazanin", font_size=12)

    doc.add_page_break()

    p_en_title = doc.add_paragraph()
    p_en_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_en_title.paragraph_format.space_before = Pt(12)
    p_en_title.paragraph_format.space_after = Pt(10)
    run_en_title = p_en_title.add_run("References (English Sources)")
    run_en_title.font.name = "Times New Roman"
    run_en_title.font.size = Pt(14)
    run_en_title.bold = True

    for ref in english_refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(ref)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    doc.save(output_docx_path)
    print(f"Exported Persian-normalized Word references to: {output_docx_path}")

def export_references_to_txt(persian_refs: List[str], english_refs: List[str], output_txt_path: str):
    with open(output_txt_path, 'w', encoding='utf-8') as f:
        f.write("منابع و مآخذ\n\n")
        f.write("الف) منابع فارسی\n")
        f.write("="*50 + "\n\n")
        for i, ref in enumerate(persian_refs, 1):
            f.write(f"{i}. {normalize_persian_text(ref)}\n\n")
        
        f.write("\n" + "="*70 + "\n")
        f.write("References (English Sources)\n")
        f.write("="*70 + "\n\n")
        for i, ref in enumerate(english_refs, 1):
            f.write(f"{ref}\n\n")
    print(f"Exported Persian-normalized text references to: {output_txt_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Compile and format thesis references.")
    parser.add_argument("--translate-dir", default=r"g:\My Drive\My Work\Fada Talebi\Translate", help="Path to Translate folder")
    parser.add_argument("--draft-docx", default=r"g:\My Drive\My Work\Fada Talebi\Thesis_Talebian.docx", help="Path to draft thesis docx")
    parser.add_argument("--output-txt", default=r"g:\My Drive\My Work\Fada Talebi\References_Compiled_APA.txt", help="Output text file")
    parser.add_argument("--output-docx", default=r"g:\My Drive\My Work\Fada Talebi\References_Compiled.docx", help="Output docx file")
    args = parser.parse_args()

    p_refs, e_refs = compile_all_references(args.translate_dir, args.draft_docx)
    export_references_to_txt(p_refs, e_refs, args.output_txt)
    export_references_to_docx(p_refs, e_refs, args.output_docx)
