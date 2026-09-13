#!/usr/bin/env python3
"""
Master Psychological & Educational Intervention Protocol Compiler
=================================================================
Author: Saber Ghaderi
Repository: GhaderiSaber/AcademicSuite
Description:
    Compiles standardized Chapter 3 APA 7 session summary tables and
    complete Appendix clinical/educational intervention manuals in Word (.docx)
    and JSON format for Master's and Doctoral theses in Psychology and Counseling.
"""

import os
import sys
import json
import argparse
from typing import Dict, List, Any, Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

FONT_TITR = "B Titr"
FONT_NAZANIN = "B Nazanin"
FONT_ENG = "Times New Roman"

# ---------------------------------------------------------------------------
# OpenXML Word Formatting Helpers (Native RTL, Borders & Fonts)
# ---------------------------------------------------------------------------
def apply_p_bidi(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(6), line_spacing=1.25):
    """Enforce RTL directionality, line spacing, and alignment on paragraph."""
    p.alignment = align
    pPr = p._element.get_or_add_pPr()
    bidi = pPr.find(qn('w:bidi'))
    if bidi is None:
        bidi = parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>')
        pPr.append(bidi)
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = line_spacing

def add_run(p, text: str, font_name: str = FONT_NAZANIN, size_pt: float = 13, bold: bool = False, italic: bool = False, color_rgb: Optional[RGBColor] = None):
    """Add run with explicit complex script font binding, w:rtl, and Latin fallback."""
    run = p.add_run(str(text))
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    rPr = run._r.get_or_add_rPr()
    sz_val = int(size_pt * 2)
    has_persian = any('\u0600' <= ch <= '\u06FF' or '\uFB50' <= ch <= '\uFDFF' or '\uFE70' <= ch <= '\uFEFF' for ch in str(text))
    if has_persian:
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} '
            f'w:ascii="{font_name}" w:hAnsi="{font_name}" '
            f'w:cs="{font_name}" w:eastAsia="{font_name}" w:hint="cs"/>'
        )
        rPr.append(rFonts)
        rtl = parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>')
        rPr.append(rtl)
    else:
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} '
            f'w:ascii="{FONT_ENG}" w:hAnsi="{FONT_ENG}" '
            f'w:cs="{font_name}" w:eastAsia="{font_name}"/>'
        )
        rPr.append(rFonts)

    szCs = parse_xml(f'<w:szCs {nsdecls("w")} w:val="{sz_val}"/>')
    rPr.append(szCs)
    if bold:
        bCs = parse_xml(f'<w:bCs {nsdecls("w")} w:val="1"/>')
        rPr.append(bCs)
    return run

def apply_table_rtl_and_borders(table, col_widths: Optional[List[float]] = None):
    """Format table with native RTL directionality and APA 7 3-line horizontal borders."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    bidiVisual = tblPr.find(qn('w:bidiVisual'))
    if bidiVisual is None:
        bidiVisual = parse_xml(f'<w:bidiVisual {nsdecls("w")}/>')
        tblPr.append(bidiVisual)

    # APA 7 borders: top line, bottom line, header underline
    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is not None:
        tblPr.remove(tblBorders)

    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

    # Underline the header row specifically
    if len(table.rows) > 0:
        header_trPr = table.rows[0]._tr.get_or_add_trPr()
        header_borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'</w:tcBorders>'
        )
        for cell in table.rows[0].cells:
            cell._tc.get_or_add_tcPr().append(parse_xml(
                f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/></w:tcBorders>'
            ))

    # Set column widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)

def add_callout_box(doc: Document, title: str, content: str, box_type: str = "metaphor"):
    """Create a shaded callout container table for clinical metaphors or experiential exercises."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    bidiVisual = tblPr.find(qn('w:bidiVisual'))
    if bidiVisual is None:
        tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))

    # Background color and border
    bg_color = "F1F5F9" if box_type == "metaphor" else "F8FAFC"
    border_color = "0284C7" if box_type == "metaphor" else "0D9488"
    icon = "💡 استعاره بالینی:" if box_type == "metaphor" else "🧘 تمرین و فن تجربی:"

    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>'))
    tcPr.append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:right w:val="single" w:sz="18" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tcBorders>'
    ))
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="120" w:type="dxa"/>'
        f'  <w:bottom w:w="120" w:type="dxa"/>'
        f'  <w:left w:w="180" w:type="dxa"/>'
        f'  <w:right w:w="180" w:type="dxa"/>'
        f'</w:tcMar>'
    ))

    # Cell paragraph for title
    p_t = cell.paragraphs[0]
    apply_p_bidi(p_t, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(4))
    add_run(p_t, f"{icon} {title}", font_name=FONT_TITR, size_pt=12, bold=True, color_rgb=RGBColor(26, 54, 93))

    # Cell paragraph for content
    p_c = cell.add_paragraph()
    apply_p_bidi(p_c, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(2))
    add_run(p_c, content, font_name=FONT_NAZANIN, size_pt=12, bold=False, color_rgb=RGBColor(30, 41, 59))

    # Spacer after callout
    p_space = doc.add_paragraph()
    apply_p_bidi(p_space, space_after=Pt(4))

def add_triad_box(doc: Document, triad: Dict[str, str]):
    """Create a high-impact callout container for the Method Triad (Motivation, Design, Advantage)."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    bidiVisual = tblPr.find(qn('w:bidiVisual'))
    if bidiVisual is None:
        tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))

    bg_color = "F5F3FF"      # Subtle violet-tinted background
    border_color = "4F46E5"  # Indigo 600

    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>'))
    tcPr.append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:right w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="DDD6FE"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="DDD6FE"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="DDD6FE"/>'
        f'</w:tcBorders>'
    ))
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="120" w:type="dxa"/>'
        f'  <w:bottom w:w="120" w:type="dxa"/>'
        f'  <w:left w:w="160" w:type="dxa"/>'
        f'  <w:right w:w="160" w:type="dxa"/>'
        f'</w:tcMar>'
    ))

    # Title paragraph
    p_t = cell.paragraphs[0]
    apply_p_bidi(p_t, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(4))
    add_run(p_t, "📐 سه‌گانه روش‌شناختی فنون جلسه (The Method Triad):", font_name=FONT_TITR, size_pt=12, bold=True, color_rgb=RGBColor(67, 56, 202))

    # Motivation
    if "motivation" in triad and triad["motivation"]:
        p_m = cell.add_paragraph()
        apply_p_bidi(p_m, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(3))
        add_run(p_m, "• چرایی و ضرورت نظری (Motivation): ", font_name=FONT_TITR, size_pt=11, bold=True, color_rgb=RGBColor(30, 41, 59))
        add_run(p_m, triad["motivation"], font_name=FONT_NAZANIN, size_pt=11.5, color_rgb=RGBColor(51, 65, 85))

    # Design
    if "design" in triad and triad["design"]:
        p_d = cell.add_paragraph()
        apply_p_bidi(p_d, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(3))
        add_run(p_d, "• طراحی و فرایند عملیاتی (Design): ", font_name=FONT_TITR, size_pt=11, bold=True, color_rgb=RGBColor(30, 41, 59))
        add_run(p_d, triad["design"], font_name=FONT_NAZANIN, size_pt=11.5, color_rgb=RGBColor(51, 65, 85))

    # Advantage
    if "advantage" in triad and triad["advantage"]:
        p_a = cell.add_paragraph()
        apply_p_bidi(p_a, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(2))
        add_run(p_a, "• مزیت رقابتی نسبت به بدیل‌ها (Advantage): ", font_name=FONT_TITR, size_pt=11, bold=True, color_rgb=RGBColor(30, 41, 59))
        add_run(p_a, triad["advantage"], font_name=FONT_NAZANIN, size_pt=11.5, color_rgb=RGBColor(51, 65, 85))

    p_space = doc.add_paragraph()
    apply_p_bidi(p_space, space_after=Pt(4))

# ---------------------------------------------------------------------------
# Document Builders
# ---------------------------------------------------------------------------
def build_protocol_docx(payload: Dict[str, Any], output_path: str):
    """Compile Word document containing Chapter 3 table and complete Appendix manual."""
    meta = payload.get("meta", {})
    sessions = payload.get("sessions", [])

    doc = Document()
    # Set standard margins (1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. Main Cover & Title
    p_title = doc.add_paragraph()
    apply_p_bidi(p_title, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(12))
    add_run(p_title, meta.get("title", "پروتکل مداخله روان‌شناختی"), font_name=FONT_TITR, size_pt=18, bold=True, color_rgb=RGBColor(26, 54, 93))

    # Sub-header Metadata
    p_meta = doc.add_paragraph()
    apply_p_bidi(p_meta, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(18))
    add_run(p_meta, f"رویکرد مداخله: {meta.get('approach', '')} • تعداد جلسات: {meta.get('session_count', len(sessions))} جلسه • مدت هر جلسه: {meta.get('session_duration_minutes', 90)} دقیقه", font_name=FONT_NAZANIN, size_pt=12, bold=True, color_rgb=RGBColor(71, 85, 105))

    # Target Population & Rationale
    p_pop = doc.add_paragraph()
    apply_p_bidi(p_pop, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(14))
    add_run(p_pop, "جامعه هدف و مشخصات اجرا: ", font_name=FONT_TITR, size_pt=13, bold=True, color_rgb=RGBColor(30, 41, 59))
    add_run(p_pop, f"این بسته درمانی برای {meta.get('target_population', 'گروه هدف')} در قالب جلسات {meta.get('format', 'گروهی')} با تواتر {meta.get('frequency', 'یک جلسه در هفته')} تدوین و بومی‌سازی شده است.", font_name=FONT_NAZANIN, size_pt=13)

    # -----------------------------------------------------------------------
    # SECTION 1: Chapter 3 APA 7 Summary Table
    # -----------------------------------------------------------------------
    p_sec1 = doc.add_paragraph()
    apply_p_bidi(p_sec1, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(6))
    add_run(p_sec1, "بخش اول: جدول خلاصه جلسات مداخله (جهت درج در فصل سوم پایان‌نامه/رساله)", font_name=FONT_TITR, size_pt=14, bold=True, color_rgb=RGBColor(26, 54, 93))

    p_tbl_cap = doc.add_paragraph()
    apply_p_bidi(p_tbl_cap, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(4))
    add_run(p_tbl_cap, "جدول ۳-X. خلاصه اهداف، محتوا و تکالیف جلسات مداخله", font_name=FONT_TITR, size_pt=12, bold=True)

    headers = ["شماره جلسه", "اهداف و محورهای اصلی جلسه", "فنون، تمرین‌ها و استعاره‌های کاربردی", "تکالیف خانگی"]
    col_widths = [1.2, 2.5, 2.3, 1.5]

    table = doc.add_table(rows=len(sessions) + 1, cols=4)
    apply_table_rtl_and_borders(table, col_widths)

    # Format Header
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        p = cell.paragraphs[0]
        apply_p_bidi(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(2))
        add_run(p, h_text, font_name=FONT_TITR, size_pt=11, bold=True)

    # Populate Rows
    for r_idx, s in enumerate(sessions):
        row_idx = r_idx + 1
        # Col 0: Session Number
        cell_0 = table.cell(row_idx, 0)
        p_0 = cell_0.paragraphs[0]
        apply_p_bidi(p_0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(2))
        add_run(p_0, f"جلسه {s.get('session_number', row_idx)}", font_name=FONT_TITR, size_pt=11, bold=True)

        # Col 1: Title & Objectives
        cell_1 = table.cell(row_idx, 1)
        p_1 = cell_1.paragraphs[0]
        apply_p_bidi(p_1, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
        add_run(p_1, f"{s.get('title', '')}\n", font_name=FONT_TITR, size_pt=10, bold=True)
        for obj in s.get("objectives", []):
            add_run(p_1, f"• {obj}\n", font_name=FONT_NAZANIN, size_pt=10)

        # Col 2: Metaphor & Experiential Technique
        cell_2 = table.cell(row_idx, 2)
        p_2 = cell_2.paragraphs[0]
        apply_p_bidi(p_2, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
        metaphor = s.get("clinical_metaphor", {})
        if metaphor:
            add_run(p_2, f"استعاره: {metaphor.get('name', '')}\n", font_name=FONT_TITR, size_pt=10, bold=True)
        tech = s.get("experiential_technique", {})
        if tech:
            add_run(p_2, f"تمرین: {tech.get('name', '')}\n", font_name=FONT_NAZANIN, size_pt=10, bold=True)
        ws = s.get("worksheet", {})
        if ws:
            add_run(p_2, f"کاربرگ: {ws.get('title', '')}\n", font_name=FONT_NAZANIN, size_pt=9, italic=True)
        triad = s.get("method_triad", {})
        if triad and triad.get("advantage"):
            add_run(p_2, f"مزیت سه‌گانه: {triad.get('advantage')}", font_name=FONT_NAZANIN, size_pt=8.5, italic=True, color_rgb=RGBColor(67, 56, 202))

        # Col 3: Homework
        cell_3 = table.cell(row_idx, 3)
        p_3 = cell_3.paragraphs[0]
        apply_p_bidi(p_3, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
        for hw in s.get("homework", []):
            add_run(p_3, f"• {hw}\n", font_name=FONT_NAZANIN, size_pt=10)

    # -----------------------------------------------------------------------
    # SECTION 2: Appendix Detailed Clinical Manual
    # -----------------------------------------------------------------------
    doc.add_page_break()
    p_sec2 = doc.add_paragraph()
    apply_p_bidi(p_sec2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(12))
    add_run(p_sec2, "بخش دوم: راهنمای تفصیلی و گام‌به‌گام جلسات مداخله (جهت درج در پیوست پایان‌نامه/رساله)", font_name=FONT_TITR, size_pt=16, bold=True, color_rgb=RGBColor(26, 54, 93))

    for s in sessions:
        # Session Heading
        p_s_head = doc.add_paragraph()
        apply_p_bidi(p_s_head, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(6))
        add_run(p_s_head, f"جلسه {s.get('session_number', '')}: {s.get('title', '')}", font_name=FONT_TITR, size_pt=14, bold=True, color_rgb=RGBColor(15, 23, 42))

        # 1. Objectives
        p_obj_head = doc.add_paragraph()
        apply_p_bidi(p_obj_head, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
        add_run(p_obj_head, "۱. اهداف ویژه جلسه:", font_name=FONT_TITR, size_pt=12, bold=True, color_rgb=RGBColor(26, 54, 93))
        for obj in s.get("objectives", []):
            p_obj = doc.add_paragraph()
            apply_p_bidi(p_obj, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
            add_run(p_obj, f"• {obj}", font_name=FONT_NAZANIN, size_pt=12)

        # 2. Psychoeducation
        p_psy_head = doc.add_paragraph()
        apply_p_bidi(p_psy_head, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
        add_run(p_psy_head, "۲. آموزش روانی و مفهوم‌بندی نظری:", font_name=FONT_TITR, size_pt=12, bold=True, color_rgb=RGBColor(26, 54, 93))
        p_psy = doc.add_paragraph()
        apply_p_bidi(p_psy, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(6))
        add_run(p_psy, s.get("psychoeducation", ""), font_name=FONT_NAZANIN, size_pt=12.5)

        # Method Triad Callout (Motivation, Design, Advantage)
        triad = s.get("method_triad", {})
        if triad:
            add_triad_box(doc, triad)

        # 3. Clinical Metaphor
        metaphor = s.get("clinical_metaphor", {})
        if metaphor:
            add_callout_box(doc, metaphor.get("name", "استعاره بالینی"), metaphor.get("description", ""), box_type="metaphor")

        # 4. Experiential Technique
        tech = s.get("experiential_technique", {})
        if tech:
            add_callout_box(doc, tech.get("name", "تمرین تجربی"), tech.get("description", ""), box_type="technique")

        # 5. Worksheet
        ws = s.get("worksheet", {})
        if ws:
            p_ws_head = doc.add_paragraph()
            apply_p_bidi(p_ws_head, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
            add_run(p_ws_head, f"۳. فعالیت کلاسی و کاربرگ: {ws.get('title', '')}", font_name=FONT_TITR, size_pt=12, bold=True, color_rgb=RGBColor(26, 54, 93))
            p_ws = doc.add_paragraph()
            apply_p_bidi(p_ws, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(6))
            add_run(p_ws, ws.get("content", ""), font_name=FONT_NAZANIN, size_pt=12)

        # 6. Homework
        p_hw_head = doc.add_paragraph()
        apply_p_bidi(p_hw_head, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
        add_run(p_hw_head, "۴. تکالیف خانگی بین‌جلسه‌ای:", font_name=FONT_TITR, size_pt=12, bold=True, color_rgb=RGBColor(26, 54, 93))
        for hw in s.get("homework", []):
            p_hw = doc.add_paragraph()
            apply_p_bidi(p_hw, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
            add_run(p_hw, f"• {hw}", font_name=FONT_NAZANIN, size_pt=12)

        # Separator line between sessions
        p_div = doc.add_paragraph()
        apply_p_bidi(p_div, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(12))
        add_run(p_div, "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━", font_name=FONT_TITR, size_pt=10, color_rgb=RGBColor(203, 213, 225))

    # 3. Primary References
    doc.add_page_break()
    p_ref_head = doc.add_paragraph()
    apply_p_bidi(p_ref_head, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(6))
    add_run(p_ref_head, "منابع استاندارد تدوین بسته درمانی/آموزشی:", font_name=FONT_TITR, size_pt=14, bold=True, color_rgb=RGBColor(26, 54, 93))
    for ref in meta.get("primary_references", []):
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_ref.paragraph_format.space_after = Pt(4)
        run_ref = p_ref.add_run(ref)
        run_ref.font.name = FONT_ENG
        run_ref.font.size = Pt(11)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"[✓] Protocol Word manual successfully compiled at: {output_path}")

def export_protocol_json(payload: Dict[str, Any], output_path: str):
    """Export clean structured protocol summary JSON for other skills."""
    meta = payload.get("meta", {})
    sessions = payload.get("sessions", [])

    table_rows = []
    for s in sessions:
        metaphor = s.get("clinical_metaphor", {}).get("name", "")
        tech = s.get("experiential_technique", {}).get("name", "")
        tech_summary = f"{metaphor} • {tech}".strip(" •")
        table_rows.append({
            "session_number": s.get("session_number"),
            "title": s.get("title"),
            "objectives": s.get("objectives", []),
            "techniques": tech_summary,
            "method_triad": s.get("method_triad", {}),
            "homework": s.get("homework", [])
        })

    summary_payload = {
        "meta": meta,
        "table_headers": ["شماره جلسه", "اهداف و محورهای اصلی", "فنون و استعاره‌ها", "تکالیف خانگی"],
        "table_rows": table_rows,
        "sessions": sessions
    }

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(summary_payload, f, ensure_ascii=False, indent=2)
    print(f"[✓] Protocol JSON summary successfully exported at: {output_path}")

def load_preset(preset_name: str, target_population: Optional[str] = None) -> Dict[str, Any]:
    """Load built-in preset or sample payload."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    sample_path = os.path.join(script_dir, "..", "examples", "sample_protocol_payload.json")

    if preset_name == "act" and os.path.exists(sample_path):
        with open(sample_path, "r", encoding="utf-8") as f:
            payload = json.load(f)
    else:
        # Generic preset template
        payload = {
            "meta": {
                "title": f"پروتکل مداخله {preset_name.upper()}",
                "approach": preset_name.upper(),
                "target_population": target_population or "جامعه هدف پژوهش",
                "session_count": 8,
                "session_duration_minutes": 90,
                "frequency": "یک جلسه در هفته",
                "format": "گروهی",
                "primary_references": ["Clinical practice manual references."]
            },
            "sessions": []
        }
        for i in range(1, 9):
            payload["sessions"].append({
                "session_number": i,
                "title": f"محور جلسه {i} بر اساس پروتکل استاندارد {preset_name.upper()}",
                "objectives": [f"هدف رفتاری و شناختی شماره یک در جلسه {i}", f"هدف رفتاری شماره دو در جلسه {i}"],
                "psychoeducation": f"مفهوم‌بندی نظری و آموزش روانی جلسه {i} بر مبنای راهنمای بالینی.",
                "clinical_metaphor": {"name": f"استعاره کاربردی جلسه {i}", "description": "تشریح استعاره ملموس بالینی جهت بازداری از مقاومت روان‌شناختی."},
                "experiential_technique": {"name": f"تکنیک درون‌جلسه‌ای {i}", "description": "تمرین تجربی و ایفای نقش مراجعان در جلسه."},
                "worksheet": {"title": f"کاربرگ شماره {i}", "content": "دستورالعمل تکمیل تمرین مکتوب درون‌جلسه‌ای."},
                "homework": [f"تکلیف رفتاری بین‌جلسه‌ای شماره یک", f"ثبت روزانه خودپایشی"]
            })

    if target_population:
        payload["meta"]["target_population"] = target_population

    return payload

def main():
    parser = argparse.ArgumentParser(description="Master Psychological & Educational Intervention Protocol Compiler")
    parser.add_argument("--json", type=str, help="Path to custom protocol payload JSON")
    parser.add_argument("--preset", type=str, default="act", choices=["act", "cbt", "schema", "cft", "mbsr", "positive", "mindful_parenting"], help="Built-in clinical preset archetype")
    parser.add_argument("--target-population", type=str, help="Target population description in Persian")
    parser.add_argument("--title", type=str, help="Custom protocol title")
    parser.add_argument("--output-docx", type=str, default="Intervention_Protocol.docx", help="Output Word document path")
    parser.add_argument("--output-json", type=str, default="protocol_summary.json", help="Output JSON summary path")
    args = parser.parse_args()

    if args.json and os.path.exists(args.json):
        with open(args.json, "r", encoding="utf-8") as f:
            payload = json.load(f)
    else:
        payload = load_preset(args.preset, args.target_population)

    if args.title:
        payload.setdefault("meta", {})["title"] = args.title
    if args.target_population:
        payload.setdefault("meta", {})["target_population"] = args.target_population

    print(f"[*] Compiling intervention protocol: '{payload.get('meta', {}).get('title', '')}'...")
    build_protocol_docx(payload, args.output_docx)
    export_protocol_json(payload, args.output_json)

if __name__ == "__main__":
    main()
