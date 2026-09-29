#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_table_alignment.py
-----------------------
Detailed empirical test of Table alignment, Cell 0 alignment, and Column 1..N centering in LibreOffice.
"""

import os
import subprocess
import xml.etree.ElementTree as ET
import docx
from docx.shared import Inches, Pt
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUTDIR = "/tmp/lo_table_test"
os.makedirs(OUTDIR, exist_ok=True)

def create_table_test_doc():
    doc = docx.Document()
    
    # sectPr bidi
    sectPr = doc.sections[0]._sectPr
    if sectPr.find(qn('w:bidi')) is None:
        sectPr.insert(0, parse_xml(f'<w:bidi {nsdecls("w")}/>'))
        
    # Table Caption variations
    # Caption 1: no jc (bidi=1)
    p1 = doc.add_paragraph()
    pPr1 = p1._p.get_or_add_pPr()
    pPr1.append(parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>'))
    r1 = p1.add_run("جدول ۱. عنوان جدول با حذف کامل w:jc (تراز طبیعی راست‌چین)")
    r1._r.get_or_add_rPr().append(parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>'))
    
    # Caption 2: jc=start
    p2 = doc.add_paragraph()
    pPr2 = p2._p.get_or_add_pPr()
    pPr2.append(parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>'))
    pPr2.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="start"/>'))
    r2 = p2.add_run("جدول ۲. عنوان جدول با w:jc=start")
    r2._r.get_or_add_rPr().append(parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>'))

    # Caption 3: jc=right
    p3 = doc.add_paragraph()
    pPr3 = p3._p.get_or_add_pPr()
    pPr3.append(parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>'))
    pPr3.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="right"/>'))
    r3 = p3.add_run("جدول ۳. عنوان جدول با w:jc=right")
    r3._r.get_or_add_rPr().append(parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>'))

    # Test Table 1: Table with bidiVisual, tblPr jc=center
    # Cell 0 has NO jc (natural RTL right)
    # Cells 1..2 have jc=center
    table = doc.add_table(rows=3, cols=3)
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))
    tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
    
    # Row 0: Headers
    headers = ["ستون اول (عنوان متغیر)", "میانگین (M)", "انحراف استاندارد (SD)"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>'))
        if j > 0:
            pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
        # If j == 0: omit jc!
        run = p.add_run(h)
        run._r.get_or_add_rPr().append(parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>'))
        
    data = [
        ["تاب‌آوری روان‌شناختی", "۲۴.۳۵", "۴.۱۸"],
        ["اضطراب فراگیر اجتماعی", "۱۸.۱۲", "۳.۷۲"]
    ]
    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            pPr = p._p.get_or_add_pPr()
            pPr.append(parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>'))
            if j > 0:
                pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
            run = p.add_run(val)
            rPr = run._r.get_or_add_rPr()
            if j == 0:
                rPr.append(parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>'))
            else:
                rPr.append(parse_xml(f'<w:rtl {nsdecls("w")} w:val="0"/>'))

    # Table Note: Justified
    p_note = doc.add_paragraph()
    pPr_n = p_note._p.get_or_add_pPr()
    pPr_n.append(parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>'))
    pPr_n.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="both"/>'))
    r_n = p_note.add_run("یادداشت. مقادیر سطح معناداری در سطح ۰.۰۵ ارزیابی شده است و متغیرها نرمال هستند.")
    r_n._r.get_or_add_rPr().append(parse_xml(f'<w:rtl {nsdecls("w")} w:val="1"/>'))
    
    docx_path = os.path.join(OUTDIR, "table_test.docx")
    doc.save(docx_path)
    return docx_path

if __name__ == '__main__':
    docx_path = create_table_test_doc()
    subprocess.run(["soffice", "--headless", "--convert-to", "fodt", docx_path, "--outdir", OUTDIR], check=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", docx_path, "--outdir", OUTDIR], check=True)

    # Parse FODT
    tree = ET.parse(os.path.join(OUTDIR, "table_test.fodt"))
    root = tree.getroot()
    ODF_NS = {
        'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
        'style': 'urn:oasis:names:tc:opendocument:xmlns:style:1.0',
        'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0',
        'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
        'fo': 'urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0',
    }

    style_props = {}
    for s in root.findall('.//style:style', ODF_NS):
        name = s.attrib.get(f"{{{ODF_NS['style']}}}name")
        props = {}
        for child in s:
            if child.tag == f"{{{ODF_NS['style']}}}paragraph-properties":
                props['text-align'] = child.attrib.get(f"{{{ODF_NS['fo']}}}text-align")
                props['writing-mode'] = child.attrib.get(f"{{{ODF_NS['style']}}}writing-mode")
            elif child.tag == f"{{{ODF_NS['style']}}}table-properties":
                props['table-align'] = child.attrib.get(f"{{{ODF_NS['table']}}}align")
                props['writing-mode'] = child.attrib.get(f"{{{ODF_NS['style']}}}writing-mode")
        style_props[name] = props

    print("=== TABLE TEST FODT PARAGRAPHS ===")
    for p in root.findall('.//text:p', ODF_NS):
        st = p.attrib.get(f"{{{ODF_NS['text']}}}style-name")
        text = "".join(p.itertext()).strip()
        if text:
            props = style_props.get(st, {})
            print(f"[{props.get('text-align', 'NONE'):^10}] [wm: {props.get('writing-mode', 'NONE'):^6}] {text}")

    print("\n=== TABLE TEST FODT TABLE PROPS ===")
    for tbl in root.findall('.//table:table', ODF_NS):
        st = tbl.attrib.get(f"{{{ODF_NS['table']}}}style-name")
        print(f"Table style: {st}, props: {style_props.get(st)}")

    # Parse PDF layout
    xml_path = os.path.join(OUTDIR, "table_test.xml")
    subprocess.run(["pdftotext", "-bbox-layout", os.path.join(OUTDIR, "table_test.pdf"), xml_path], check=True)
    pdf_tree = ET.parse(xml_path)
    pdf_ns = {'h': 'http://www.w3.org/1999/xhtml'}
    print("\n=== TABLE TEST PDF PHYSICAL LAYOUT ===")
    for line in pdf_tree.findall('.//h:line', pdf_ns):
        words = [w.text for w in line.findall('.//h:word', pdf_ns) if w.text]
        line_text = ' '.join(words)
        xMin = float(line.attrib['xMin'])
        xMax = float(line.attrib['xMax'])
        pos = 'RIGHT' if xMax > 500 else ('LEFT' if xMin < 100 else 'MIDDLE')
        print(f"[{pos:6}] (xMin={xMin:5.1f}, xMax={xMax:5.1f}) | {line_text}")

