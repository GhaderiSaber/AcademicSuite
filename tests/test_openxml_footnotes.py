import os
import tempfile
import json
import pytest
import docx
import zipfile
import xml.etree.ElementTree as ET

import sys
sys.path.insert(0, os.path.abspath(".agents/skills/apa-reporting/scripts"))
from inject_openxml_footnotes import OpenXMLFootnoteInjector


def test_extract_footnotes_from_markdown(tmp_path):
    md_content = """# Title
Some Persian narrative text with footnote[^1] and another[^2].

## References
[^1]: Hackman & Oldham (1976)
[^2]: Demircioglu (2020)
"""
    md_file = tmp_path / "test.md"
    md_file.write_text(md_content, encoding="utf-8")

    footnotes = OpenXMLFootnoteInjector.extract_footnotes_from_markdown(str(md_file))
    assert len(footnotes) == 2
    assert footnotes[0] == (1, "Hackman & Oldham (1976)")
    assert footnotes[1] == (2, "Demircioglu (2020)")


def test_extract_footnotes_from_json(tmp_path):
    # Test format A: list of dicts
    json_data_a = [
        {"id": 1, "text": "First Author"},
        {"id": 2, "text": "Second Author"}
    ]
    json_file_a = tmp_path / "fn_a.json"
    json_file_a.write_text(json.dumps(json_data_a), encoding="utf-8")

    fns_a = OpenXMLFootnoteInjector.extract_footnotes_from_json(str(json_file_a))
    assert fns_a == [(1, "First Author"), (2, "Second Author")]

    # Test format B: dict
    json_data_b = {
        "1": "First Author",
        "2": "Second Author"
    }
    json_file_b = tmp_path / "fn_b.json"
    json_file_b.write_text(json.dumps(json_data_b), encoding="utf-8")

    fns_b = OpenXMLFootnoteInjector.extract_footnotes_from_json(str(json_file_b))
    assert fns_b == [(1, "First Author"), (2, "Second Author")]


def test_inject_footnotes_end_to_end(tmp_path):
    # 1. Create a dummy docx with python-docx
    doc_path = tmp_path / "sample.docx"
    doc = docx.Document()
    p1 = doc.add_paragraph("این یک متن نمونه فارسی است که نام هکمن و اولدهام[^1] در آن ذکر شده است.")
    p2 = doc.add_paragraph("همچنین نظریه بار شناختی[^2] نیز در اینجا مطرح گردیده است.")
    p_trailing_1 = doc.add_paragraph("[^1]: Hackman & Oldham")
    p_trailing_2 = doc.add_paragraph("[^2]: Cognitive Load Theory")
    doc.save(str(doc_path))

    # Create dummy metadata json
    meta_path = tmp_path / "sample.json"
    meta_data = {"stage": "test_stage", "hypothesis": 1}
    meta_path.write_text(json.dumps(meta_data), encoding="utf-8")

    footnotes = [
        (1, "Hackman & Oldham"),
        (2, "Cognitive Load Theory")
    ]

    out_docx = tmp_path / "sample_out.docx"
    result = OpenXMLFootnoteInjector.inject(
        docx_path=str(doc_path),
        footnotes=footnotes,
        output_docx_path=str(out_docx),
        update_json_path=str(meta_path)
    )

    assert result["success"] is True
    assert result["footnotes_injected"] == 2
    assert result["remaining_paragraphs"] == 2  # Trailing 2 paragraphs removed, 2 narrative remain

    # Verify JSON was updated
    with open(str(meta_path), "r", encoding="utf-8") as f:
        updated_meta = json.load(f)
    assert "docx_sha256" in updated_meta
    assert updated_meta["footnote_count"] == 2

    # Forensic verification
    v_report = OpenXMLFootnoteInjector.verify(str(out_docx), expected_footnote_count=2)
    assert v_report["valid"] is True
    assert v_report["paragraph_count"] == 2
    assert v_report["has_footnotes_xml"] is True
    assert v_report["has_content_type_override"] is True
    assert v_report["has_document_relationship"] is True
    assert v_report["footnote_count"] == 2
    assert v_report["references_match_definitions"] is True

    # Ensure python-docx can still open and parse it without error
    reloaded_doc = docx.Document(str(out_docx))
    assert len(reloaded_doc.paragraphs) == 2
    # Verify markers were replaced (no raw [^1] or [^2] in text)
    assert "[^1]" not in reloaded_doc.paragraphs[0].text
    assert "[^2]" not in reloaded_doc.paragraphs[1].text


def test_cli_invocation(tmp_path):
    import subprocess
    # Create sample docx
    doc_path = tmp_path / "cli_sample.docx"
    doc = docx.Document()
    doc.add_paragraph("متن تستی با ارجاع به منبع[^1].")
    doc.save(str(doc_path))

    # Create companion markdown draft
    md_path = tmp_path / "cli_sample.md"
    md_path.write_text("متن تستی با ارجاع به منبع[^1].\n\n[^1]: Author (2024)\n", encoding="utf-8")

    out_docx = tmp_path / "cli_sample_out.docx"
    report_json = tmp_path / "report.json"

    script_path = os.path.abspath(".agents/skills/apa-reporting/scripts/inject_openxml_footnotes.py")
    cmd = [
        sys.executable,
        script_path,
        "--docx", str(doc_path),
        "--output", str(out_docx),
        "--md", str(md_path),
        "--verify",
        "--report-path", str(report_json)
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    assert res.returncode == 0
    assert "Footnotes injected successfully" in res.stdout
    assert os.path.exists(str(out_docx))
    assert os.path.exists(str(report_json))

    with open(str(report_json), "r", encoding="utf-8") as f:
        rep = json.load(f)
    assert rep["valid"] is True
    assert rep["footnote_count"] == 1


def test_error_handling(tmp_path):
    with pytest.raises(FileNotFoundError):
        OpenXMLFootnoteInjector.extract_footnotes_from_markdown("non_existent_file.md")

    with pytest.raises(FileNotFoundError):
        OpenXMLFootnoteInjector.extract_footnotes_from_json("non_existent_file.json")

    with pytest.raises(FileNotFoundError):
        OpenXMLFootnoteInjector.inject("non_existent.docx", [(1, "Test")])

    doc_path = tmp_path / "empty_doc.docx"
    doc = docx.Document()
    doc.add_paragraph("Paragraph")
    doc.save(str(doc_path))

    with pytest.raises(ValueError):
        OpenXMLFootnoteInjector.inject(str(doc_path), [])

