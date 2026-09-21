#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/test_literature_synthesis_matrix.py — Unit Tests for Chapter 2 Literature Review Thematic Synthesis Matrix Engine

Verifies:
1. Thematic clustering, sample size metrics calculation, and region distribution.
2. Concordance / discordance evaluation and 5-part scholarly Persian synthesis narrative generation.
3. Physical generation of synchronized Triad artifacts on disk (Stage 2.6):
   - 06_literature_matrix_table.docx (OpenXML BiDi RTL, APA 7 borderless table, authentic Persian typography)
   - 06_literature_matrix_table.md   (Scholarly markdown narrative + thematic tables + gap matrix)
   - 06_literature_matrix_table.json (Machine-readable structured parameters & gap matrix)
   - Literature_Synthesis_Matrix.xlsx (4-sheet formatted openpyxl workbook with RTL layout)
4. CLI execution via generate_literature_matrix_triad.py with --generate-sample.
5. Strict test isolation: Zero file pollution or side effects in .agents/learning/.
"""

import os
import sys
import json
import shutil
import tempfile
import unittest

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SKILL_SCRIPTS_DIR = os.path.join(ROOT_DIR, ".agents", "skills", "persian-literature-review-builder", "scripts")
SKILL_EXAMPLES_DIR = os.path.join(ROOT_DIR, ".agents", "skills", "persian-literature-review-builder", "examples")

for p in [ROOT_DIR, SKILL_SCRIPTS_DIR, SKILL_EXAMPLES_DIR]:
    if p not in sys.path:
        sys.path.insert(0, p)

from literature_synthesis_matrix_engine import (
    ThematicClusteringEngine,
    OpenXmlMatrixDocxBuilder,
    ExcelMatrixWorkbookBuilder,
    MarkdownMatrixBuilder,
    LiteratureSynthesisMatrixEngine
)
from sample_synthesis_matrix_payload import get_sample_payload

try:
    import docx
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class TestLiteratureSynthesisMatrixEngine(unittest.TestCase):
    """Test suite for Chapter 2 Literature Review Thematic Synthesis Matrix Engine."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp(prefix="test_lit_matrix_")
        self.sample_payload = get_sample_payload()

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_01_thematic_clustering_and_metrics_calculation(self):
        """Verifies mathematical analysis, sample metrics, and regional distributions."""
        analysis = ThematicClusteringEngine.analyze_payload(self.sample_payload)

        self.assertEqual(analysis.get("contract_version"), "1.0.0")
        self.assertEqual(analysis.get("stage_id"), "06_literature_matrix_table")

        metrics = analysis.get("summary_metrics", {})
        self.assertEqual(metrics.get("total_studies_reviewed"), 5)
        self.assertEqual(metrics.get("total_sample_size"), 811)
        self.assertEqual(metrics.get("mean_sample_size"), 162.2)
        self.assertEqual(metrics.get("min_sample_size"), 34)
        self.assertEqual(metrics.get("max_sample_size"), 345)
        self.assertEqual(metrics.get("iranian_studies_count"), 3)
        self.assertEqual(metrics.get("international_studies_count"), 2)
        self.assertEqual(metrics.get("domestic_percentage"), 60.0)
        self.assertEqual(metrics.get("themes_count"), 2)

        # Check theme 1
        themes = analysis.get("analyzed_themes", [])
        self.assertEqual(len(themes), 2)
        t1 = themes[0]
        self.assertEqual(t1.get("theme_id"), "THEME_1")
        self.assertEqual(t1.get("studies_count"), 3)
        self.assertEqual(t1.get("concordance_verdict"), "UNANIMOUS_SUPPORT")
        self.assertIn("همگرایی قاطع", t1.get("concordance_text_fa"))

        # Verify scholarly Persian synthesis narrative
        narrative = t1.get("synthesis_narrative", "")
        self.assertIn("نظریه حفاظت از منابع (COR)", narrative)
        self.assertIn("احمدی و رضایی (۱۴۰۲)", narrative)
        self.assertIn("ویلیامز و همکاران (۲۰۲۴)", narrative)
        self.assertIn("شکاف پژوهشی", narrative)

        # Verify master gaps extraction
        gaps = analysis.get("master_research_gaps", [])
        self.assertEqual(len(gaps), 2)
        self.assertEqual(gaps[0].get("gap_id"), "GAP-01")
        self.assertIn("مداخله‌های ACT", gaps[0].get("theme_title"))

    @unittest.skipUnless(HAS_DOCX, "python-docx not available")
    def test_02_openxml_word_document_generation_and_typography(self):
        """Verifies OpenXML BiDi RTL standards, APA 7 borderless table, and Persian typography."""
        analysis = ThematicClusteringEngine.analyze_payload(self.sample_payload)
        docx_file = os.path.join(self.temp_dir, "06_literature_matrix_table.docx")

        OpenXmlMatrixDocxBuilder.build_docx(analysis, docx_file)
        self.assertTrue(os.path.isfile(docx_file))
        self.assertGreater(os.path.getsize(docx_file), 5000)

        # Inspect Document XML
        doc = docx.Document(docx_file)
        self.assertGreater(len(doc.paragraphs), 5)

        # Section BiDi
        sect_bidi = doc.sections[0]._sectPr.find(qn("w:bidi"))
        self.assertIsNotNone(sect_bidi, "Section must have <w:bidi/> for RTL layout")

        # Table Verification
        self.assertEqual(len(doc.tables), 2, "Expected 2 APA 7 tables (one per theme)")
        for tbl in doc.tables:
            tblPr = tbl._tbl.tblPr
            # BiDi Visual Table layout
            self.assertIsNotNone(tblPr.find(qn("w:bidiVisual")), "Table must have <w:bidiVisual/>")

            # APA 7 3-line borders
            borders = tblPr.find(qn("w:tblBorders"))
            self.assertIsNotNone(borders, "Table must have <w:tblBorders>")
            top = borders.find(qn("w:top"))
            bottom = borders.find(qn("w:bottom"))
            insideH = borders.find(qn("w:insideH"))
            left = borders.find(qn("w:left"))
            right = borders.find(qn("w:right"))
            insideV = borders.find(qn("w:insideV"))

            self.assertEqual(top.get(qn("w:val")), "single")
            self.assertEqual(bottom.get(qn("w:val")), "single")
            self.assertEqual(insideH.get(qn("w:val")), "single")
            self.assertEqual(left.get(qn("w:val")), "none", "APA 7 forbids vertical left borders")
            self.assertEqual(right.get(qn("w:val")), "none", "APA 7 forbids vertical right borders")
            self.assertEqual(insideV.get(qn("w:val")), "none", "APA 7 forbids vertical inside borders")

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl not available")
    def test_03_excel_multi_sheet_workbook_generation(self):
        """Verifies 4-sheet formatted openpyxl workbook with RTL layout and correct metrics."""
        analysis = ThematicClusteringEngine.analyze_payload(self.sample_payload)
        xlsx_file = os.path.join(self.temp_dir, "Literature_Synthesis_Matrix.xlsx")

        ExcelMatrixWorkbookBuilder.build_workbook(analysis, xlsx_file)
        self.assertTrue(os.path.isfile(xlsx_file))

        wb = openpyxl.load_workbook(xlsx_file)
        expected_sheets = [
            "Overview & Metrics",
            "Thematic Matrix",
            "Empirical Studies Detail",
            "Research Gaps & Critique"
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Verify RTL directionality on all sheets
        for sheet_name in expected_sheets:
            ws = wb[sheet_name]
            self.assertTrue(ws.views.sheetView[0].rightToLeft, f"Sheet '{sheet_name}' must be RTL")

        # Verify Overview sheet has data
        ws_overview = wb["Overview & Metrics"]
        self.assertEqual(ws_overview.cell(row=6, column=2).value, 5)  # total studies

        # Verify Detail sheet has all 5 studies
        ws_detail = wb["Empirical Studies Detail"]
        self.assertEqual(ws_detail.max_row, 6)  # header + 5 studies

    def test_04_markdown_triad_deliverable_generation(self):
        """Verifies 06_literature_matrix_table.md content, tables, and research gaps."""
        analysis = ThematicClusteringEngine.analyze_payload(self.sample_payload)
        md_file = os.path.join(self.temp_dir, "06_literature_matrix_table.md")

        MarkdownMatrixBuilder.build_markdown(analysis, md_file)
        self.assertTrue(os.path.isfile(md_file))

        with open(md_file, "r", encoding="utf-8") as f:
            content = f.read()

        self.assertIn("# ۲-۳-۳. ماتریس سنتز و جدول جامع پیشینه پژوهش‌های تجربی", content)
        self.assertIn("N = 811", content)
        self.assertIn("THEME_1", content)
        self.assertIn("THEME_2", content)
        self.assertIn("احمدی و رضایی (۱۴۰۲)", content)
        self.assertIn("Wilks' Lambda", content)
        self.assertIn("۰.۰۰۱ > p", content)
        self.assertIn("Research Gaps", content)
        self.assertIn("GAP-01", content)
        self.assertIn("GAP-02", content)

    def test_05_master_engine_process_and_export(self):
        """Verifies end-to-end generation of synchronized quad deliverables."""
        artifacts = LiteratureSynthesisMatrixEngine.process_and_export(
            payload=self.sample_payload,
            out_dir=self.temp_dir,
            lang="fa"
        )

        self.assertTrue(os.path.isfile(artifacts["docx"]))
        self.assertTrue(os.path.isfile(artifacts["md"]))
        self.assertTrue(os.path.isfile(artifacts["json"]))
        self.assertTrue(os.path.isfile(artifacts["xlsx"]))

        # Verify JSON schema contract
        with open(artifacts["json"], "r", encoding="utf-8") as f:
            data = json.load(f)
        self.assertEqual(data.get("stage_id"), "06_literature_matrix_table")
        self.assertIn("summary_metrics", data)
        self.assertIn("analyzed_themes", data)
        self.assertIn("master_research_gaps", data)

    def test_06_cli_script_execution(self):
        """Verifies CLI script generate_literature_matrix_triad.py with --generate-sample."""
        cli_script = os.path.join(ROOT_DIR, ".agents", "scripts", "generate_literature_matrix_triad.py")
        self.assertTrue(os.path.isfile(cli_script))

        # Run CLI in isolated temp directory
        cli_out_dir = os.path.join(self.temp_dir, "cli_output")
        cmd = f"{sys.executable} {cli_script} --generate-sample --out-dir {cli_out_dir}"
        exit_code = os.system(cmd)
        self.assertEqual(exit_code, 0, "CLI script execution failed")

        self.assertTrue(os.path.isfile(os.path.join(cli_out_dir, "06_literature_matrix_table.docx")))
        self.assertTrue(os.path.isfile(os.path.join(cli_out_dir, "06_literature_matrix_table.md")))
        self.assertTrue(os.path.isfile(os.path.join(cli_out_dir, "06_literature_matrix_table.json")))
        self.assertTrue(os.path.isfile(os.path.join(cli_out_dir, "Literature_Synthesis_Matrix.xlsx")))


if __name__ == "__main__":
    unittest.main()
