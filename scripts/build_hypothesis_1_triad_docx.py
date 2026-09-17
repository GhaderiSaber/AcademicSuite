#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_hypothesis_1_triad_docx.py — Compiles OpenXML Word deliverable for Hypothesis 1.
"""
import os
import sys

# Virtualenv auto-discovery shim
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
for venv_name in [".venv", "venv"]:
    venv_lib = os.path.join(ROOT_DIR, venv_name, "lib")
    if os.path.isdir(venv_lib):
        for entry in os.listdir(venv_lib):
            sp = os.path.join(venv_lib, entry, "site-packages")
            if os.path.isdir(sp) and sp not in sys.path:
                sys.path.insert(0, sp)

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_p_rtl(p, justify=True):
    pPr = p._element.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_p(doc, text, font_name="B Nazanin", size=13, bold=False, italic=False):
    p = doc.add_paragraph()
    set_p_rtl(p, justify=True)
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    rPr = r._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_p_rtl(p, justify=False)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    font_name = "B Titr"
    size = 14 if level == 1 else 13
    r.font.name = font_name
    r.font.size = Pt(size)
    r.bold = True
    rPr = r._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)
    return p


def build_docx(out_path):
    doc = Document()

    add_heading(doc, "آزمون فرضیه اول: پیش‌بینی فرسودگی شغلی بر اساس استرس شغلی و انعطاف‌پذیری روان‌شناختی", level=1)

    add_heading(doc, "۱. بیان فرضیه و صورت‌بندی مدل", level=2)
    add_p(doc, "فرضیه اول پژوهش بیان می‌دارد که «استرس شغلی و انعطاف‌پذیری روان‌شناختی به طور همزمان فرسودگی شغلی کارکنان را پیش‌بینی می‌کنند». به منظور ارزیابی سهم نسبی هر یک از متغیرهای پیش‌بین در تبیین واریانس متغیر ملاک، از الگوی رگرسیون خطی چندگانه به روش همزمان (Enter) استفاده شد.")

    add_heading(doc, "۲. بررسی مفروضه‌های آماری رگرسیون", level=2)
    add_p(doc, "پیش از اجرای تحلیل رگرسیون، پیش‌شرط‌های اساسی شامل نرمال بودن توزیع خطاهای برآورد، خطی بودن رابطه متغیرها و عدم هم‌خطی چندگانه مورد ارزیابی قرار گرفت. شاخص‌های چولگی و کشیدگی متغیرها در محدوده استاندارد (بین ۱- تا ۱+) قرار داشت که نشان‌دهنده برقراری فرض توزیع نرمال است. همچنین، شاخص عامل تورم واریانس (VIF) برای هر دو متغیر پیش‌بین برابر با ۱.۰۹ و شاخص رواداری (Tolerance) برابر با ۰.۹۲ محاسبه شد که حاکی از فقدان هم‌خطی چندگانه بحرانی میان متغیرهای پیش‌بین است.")

    add_heading(doc, "۳. یافته‌های استنباطی رگرسیون چندگانه", level=2)
    add_p(doc, "خلاصه مشخصه‌های آماری مدل رگرسیون و ضرایب استاندارد و غیراستاندارد متغیرهای پیش‌بین در جدول ۱ گزارش شده است.")

    # Table
    table_title = doc.add_paragraph()
    set_p_rtl(table_title, justify=False)
    table_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_title = table_title.add_run("جدول ۱. ضرایب رگرسیون چندگانه پیش‌بینی فرسودگی شغلی بر اساس استرس شغلی و انعطاف‌پذیری روان‌شناختی (۱۰۰ = N)")
    r_title.font.name = "B Nazanin"
    r_title.bold = True
    r_title.font.size = Pt(11)

    tbl = doc.add_table(rows=4, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    bidiVisual = OxmlElement('w:bidiVisual')
    tblPr.append(bidiVisual)

    headers = ["متغیر پیش‌بین", "B", "SE", "β", "t", "سطح معناداری", "VIF"]
    rows_data = [
        ["مقدار ثابت", "۲.۳۱۴", "۰.۳۵۲", "—", "۶.۵۷۴", "۰.۰۰۱ > p", "—"],
        ["استرس شغلی", "۰.۴۶۶", "۰.۰۸۱", "۰.۴۶۴", "۵.۷۵۶", "۰.۰۰۱ > p", "۱.۰۹"],
        ["انعطاف‌پذیری روان‌شناختی", "۰.۳۶۲-", "۰.۰۸۶", "۰.۳۴۰-", "۴.۲۲۴-", "۰.۰۰۱ > p", "۱.۰۹"]
    ]

    for c_idx, h in enumerate(headers):
        cell = tbl.rows[0].cells[c_idx]
        p = cell.paragraphs[0]
        set_p_rtl(p, justify=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "B Nazanin"
        r.bold = True
        r.font.size = Pt(11)

    for r_idx, row_values in enumerate(rows_data):
        for c_idx, val in enumerate(row_values):
            cell = tbl.rows[r_idx + 1].cells[c_idx]
            p = cell.paragraphs[0]
            set_p_rtl(p, justify=False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Times New Roman" if any(char.isdigit() or char in ["-", ">", "p"] for char in val) else "B Nazanin"
            r.font.size = Pt(11)

    note_p = doc.add_paragraph()
    set_p_rtl(note_p, justify=True)
    r_note = note_p.add_run("یادداشت. متغیر ملاک: فرسودگی شغلی. شاخص‌های برازش الگو: 34.389 = (۹۷ ،۲)F، 0.415 = R²، 0.403 = تعدیل‌شده R²، ۰.۰۰۱ > p.")
    r_note.font.name = "B Nazanin"
    r_note.font.size = Pt(10)
    r_note.italic = True

    add_heading(doc, "۴. تحلیل و تفسیر آماری", level=2)
    add_p(doc, "بر اساس یافته‌های حاصل از آزمون تحلیل واریانس رگرسیون، مدل رگرسیونی ترسیم‌شده با اطمینان ۹۹ درصد معنادار است (34.389 = (۹۷ ،۲)F، ۰.۰۰۱ > p). ضریب تعیین محاسبه‌شده نشان می‌دهد که ۴۱.۵ درصد از کل تغییرات و واریانس فرسودگی شغلی کارکنان (0.415 = R²) توسط متغیرهای استرس شغلی و انعطاف‌پذیری روان‌شناختی تبیین می‌گردد.")
    add_p(doc, "بررسی ضرایب استاندارد رگرسیون (β) نشان می‌دهد که متغیر استرس شغلی با ضریب استاندارد ۰.۴۶۴ و مقدار آماره ۵.۷۵۶ = t در سطح ۰.۰۰۱ > p به صورت مثبت و معنادار قادر به پیش‌بینی فرسودگی شغلی است؛ بدین معنا که با افزایش یک انحراف استاندارد در میزان استرس شغلی، فرسودگی شغلی به میزان ۰.۴۶۴ انحراف استاندارد افزایش می‌یابد. همچنین، متغیر انعطاف‌پذیری روان‌شناختی با ضریب استاندارد ۰.۳۴۰- و مقدار آماره ۴.۲۲۴- = t در سطح ۰.۰۰۱ > p به صورت منفی و معنادار فرسودگی شغلی را پیش‌بینی می‌نماید؛ بدین ترتیب، کارکنانی که از سطوح بالاتری از پذیرش و انعطاف‌پذیری در مواجهه با تجارب درونی دشوار برخوردارند، نشانه‌های فرسودگی شغلی کمتری را تجربه می‌کنند.")

    add_heading(doc, "۵. نتیجه‌گیری آماری", level=2)
    add_p(doc, "با عنایت به معناداری کلی مدل رگرسیونی و معناداری انفرادی ضرایب هر دو متغیر در جهت مورد انتظار، فرضیه اول پژوهش تأیید می‌گردد.")

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    print(f"Saved DOCX artifact: {out_path}")


if __name__ == "__main__":
    out_docx = sys.argv[1] if len(sys.argv) > 1 else "06_hypothesis_1_regression.docx"
    build_docx(out_docx)
